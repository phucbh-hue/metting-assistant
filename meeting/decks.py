"""Thư viện tài liệu trên máy: tìm và mở PDF, PowerPoint, Word, Markdown để trình bày trong cuộc họp.

"Jarvis ơi, mở file báo cáo Q3 trong Downloads" -> tìm theo tên thư mục / tên tệp trong các thư mục tài liệu
(thư mục slides/ của dự án, Desktop, Documents, Downloads, OneDrive, cùng các thư mục thêm trong Cài đặt hoặc
LIBRARY_DIRS), đọc thành bộ slide của hệ thống.

Định dạng:
- .pdf  : mỗi trang thành một slide hiện đúng hình trang gốc (dựng ảnh bằng pypdfium2), kèm chữ của trang để tìm kiếm,
          tự chuyển slide theo lời nói và viết lời thuyết trình.
- .pptx : hiện ĐÚNG thiết kế từng slide: LibreOffice chuyển sang PDF rồi dựng ảnh từng trang (máy chủ bản web cài sẵn
          LibreOffice và font hay gặp). Chữ và ghi chú người trình bày (kịch bản) vẫn đọc từ tệp gốc; slide ẩn bỏ qua
          như khi trình chiếu. Bộ slide toàn ảnh (xuất từ Canva...) dùng thẳng ảnh gốc. Máy không có LibreOffice thì
          dựng lại theo bố cục của ứng dụng (chữ + hình lớn nhất), source["exact"] = False để trợ lý báo người dùng.
          Ảnh dựng xong giữ theo nội dung tệp: mở lại đúng tệp đó (kể cả gửi lên lại) thì không phải dựng lại.
- .ppt  : cần LibreOffice để chuyển đổi.
- .docx : tiêu đề lớn thành slide, đoạn văn / gạch đầu dòng thành ý.
- .md / .txt : "#" là tên bộ slide, mỗi "##" một slide, gạch đầu dòng là ý, đoạn văn là ghi chú.
- .json : đúng cấu trúc deck của hệ thống.

Chỉ tên tệp được quét; nội dung chỉ được đọc khi người dùng chọn mở tệp đó.
"""
import hashlib
import json
import logging
import os
import re
import shutil
import signal
import subprocess
import tempfile
import threading
import time
import unicodedata
from pathlib import Path
from typing import Any, Dict, List, Optional

ROOT = Path(__file__).resolve().parent.parent
SLIDES_DIR = Path(os.getenv("SLIDES_DIR") or (ROOT / "slides"))
ASSETS_DIR = Path(os.getenv("DECK_ASSETS_DIR") or (ROOT / "data" / "deck_assets"))
EXTS = (".pdf", ".pptx", ".ppt", ".docx", ".json", ".md", ".txt")
MAX_PAGES = int(os.getenv("DECK_MAX_PAGES", "60") or 60)       # trang tối đa đọc từ một tệp
SCAN_DEPTH = 4                                                  # độ sâu thư mục khi quét tên tệp
FULL_BLEED = 0.9                                                # hình PowerPoint phủ từ 90% slide: hiện cả hình
PURE_IMAGE = 0.98                                               # slide chỉ có một hình phủ kín, không chữ: đúng thiết kế
PAGE_PX = 2000                                                  # cạnh dài ảnh mỗi trang / slide (điểm ảnh)
RENDER_VERSION = "2"                                            # đổi cách dựng ảnh thì đổi số này (ảnh cũ không dùng lại)
LO_TIMEOUT_S = 240
LO_PROFILE = ROOT / "data" / "lo-profile"                       # hồ sơ LibreOffice riêng của Copilot
_LO_LOCK = threading.Lock()                                     # LibreOffice: mỗi lúc một tệp (giới hạn RAM, CPU)
log = logging.getLogger("meeting.decks")
SCAN_LIMIT = 4000
_SKIP_DIRS = {"node_modules", ".git", "__pycache__", "appdata", "$recycle.bin", "venv", ".venv", "site-packages",
              "program files", "program files (x86)", "windows", ".cache", ".npm", ".vscode", "data"}
_STOP = {"slide", "slides", "bo", "bai", "file", "tep", "folder", "thu", "muc", "o", "trong", "tu", "cua", "cai",
         "giup", "anh", "chi", "em", "minh", "nhe", "nha", "di", "cho", "xem", "mo", "ra", "len", "ve", "ten", "la",
         "tai", "lieu", "trinh", "bay", "hay", "cai", "nay", "do", "va", "voi"}
# Tên thư mục người dùng hay gọi -> từ khóa của thư mục gốc
_ROOT_WORDS = {"downloads": {"download", "downloads", "xuong", "ve"}, "desktop": {"desktop", "man", "hinh"},
               "documents": {"documents", "document", "docs"}}
_EXT_WORDS = {".pdf": {"pdf"}, ".pptx": {"powerpoint", "ppt", "pptx", "slide"}, ".ppt": {"powerpoint", "ppt"},
              ".docx": {"word", "docx", "doc"}, ".md": {"md", "markdown"}}


def fold(text: str) -> str:
    """Bỏ dấu tiếng Việt, chữ thường: "Mega Sale 10.10" -> "mega sale 10.10"."""
    t = unicodedata.normalize("NFD", str(text or "")).replace("\u0111", "d").replace("\u0110", "d")
    return "".join(c for c in t if unicodedata.category(c) != "Mn").lower()


def _tokens(text: str) -> List[str]:
    return [w for w in re.findall(r"[a-z0-9]+", fold(text)) if w not in _STOP]


# ------------------------------------------------------------ thư mục tài liệu ---
def _default_dirs() -> List[Path]:
    if os.getenv("LIBRARY_DEFAULTS", "1").strip() == "0":
        return []
    home = Path.home()
    out = [home / "Desktop", home / "Documents", home / "Downloads"]
    for od in sorted(home.glob("OneDrive*")):
        out += [od / "Desktop", od / "Documents"]
    return out


def _custom_dirs() -> List[Path]:
    dirs: List[str] = [d for d in (os.getenv("LIBRARY_DIRS") or "").split(os.pathsep) if d.strip()]
    try:
        from meeting import db
        saved = json.loads(db.get_setting("library_dirs", "[]") or "[]")
        dirs += [str(d) for d in saved if str(d).strip()]
    except Exception:
        pass
    return [Path(os.path.expandvars(os.path.expanduser(d.strip()))) for d in dirs]


def library_roots() -> List[Path]:
    """Các thư mục trợ lý được phép tìm tệp, theo thứ tự ưu tiên, bỏ trùng và thư mục không tồn tại."""
    seen, out = set(), []
    for p in [SLIDES_DIR] + _custom_dirs() + _default_dirs():
        try:
            key = str(p.resolve()).lower()
        except Exception:
            continue
        if key in seen or not p.is_dir():
            continue
        seen.add(key)
        out.append(p)
    return out


def _walk(base: Path, depth: int = SCAN_DEPTH, limit: int = SCAN_LIMIT) -> List[Path]:
    out: List[Path] = []
    base_depth = len(base.parts)
    for cur, dirs, files in os.walk(base):
        cp = Path(cur)
        dirs[:] = [d for d in dirs if not d.startswith((".", "~", "$")) and d.lower() not in _SKIP_DIRS]
        if len(cp.parts) - base_depth >= depth:
            dirs[:] = []
        for f in files:
            if f.startswith((".", "~$")) or f.lower() == "readme.md":
                continue
            if os.path.splitext(f)[1].lower() in EXTS:
                out.append(cp / f)
                if len(out) >= limit:
                    return out
    return out


_CACHE: Dict[str, Any] = {"at": 0.0, "key": "", "files": []}


def list_files(base: Optional[Path] = None) -> List[Path]:
    """Tệp đọc được trong một thư mục (base) hoặc trong mọi thư mục tài liệu (quét tên, lưu tạm 30 giây)."""
    if base is not None:
        return sorted(_walk(Path(base)))[:500] if Path(base).is_dir() else []
    roots = library_roots()
    key = "|".join(str(r) for r in roots)
    if _CACHE["key"] == key and time.time() - _CACHE["at"] < 30:
        return list(_CACHE["files"])
    files: List[Path] = []
    for r in roots:
        files += _walk(r, limit=max(0, SCAN_LIMIT - len(files)))
        if len(files) >= SCAN_LIMIT:
            break
    _CACHE.update(at=time.time(), key=key, files=files)
    return list(files)


def invalidate() -> None:
    """Đổi danh sách thư mục: quét lại tên tệp ở lần tìm sau."""
    _CACHE.update(at=0.0, key="", files=[])


def allowed(path: str) -> Optional[Path]:
    """Tệp được phép mở: đúng định dạng và nằm trong một thư mục tài liệu (không mở tệp bất kỳ trên máy)."""
    try:
        p = Path(path).resolve()
    except (OSError, ValueError):
        return None
    if p.suffix.lower() not in EXTS or not p.is_file():
        return None
    for r in library_roots():
        try:
            p.relative_to(r.resolve())
            return p
        except ValueError:
            continue
    return None


def _root_of(p: Path, roots: List[Path]) -> Optional[Path]:
    for r in roots:
        try:
            p.relative_to(r)
            return r
        except ValueError:
            continue
    return None


_PATH_RE = re.compile(r"([A-Za-z]:[\\/][^\"<>|?*\n]+?\.(?:pdf|pptx|ppt|docx|md|txt|json))", re.I)


def find_files(query: str, base: Optional[Path] = None, exclude: Optional[str] = None) -> List[Dict[str, Any]]:
    """Xếp hạng tệp theo mức khớp giữa câu nói và tên thư mục / tên tệp (có thể gõ thẳng đường dẫn đầy đủ)."""
    m = _PATH_RE.search(query or "")
    if m and allowed(m.group(1)):
        p = allowed(m.group(1))
        return [{"path": str(p), "rel": p.name, "folder": p.parent.name, "name": p.stem, "ext": p.suffix.lower(),
                 "root": str(p.parent), "score": 99.0}]
    roots = [Path(base)] if base is not None else library_roots()
    files = list_files(base) if base is not None else list_files()
    q = set(_tokens(query))
    fq = fold(query)
    results = []
    for p in files:
        if exclude and os.path.normcase(str(p)) == os.path.normcase(str(exclude)):
            continue
        root = _root_of(p, roots) or p.parent
        rel = p.relative_to(root)
        score = _score(q, fq, rel.parts[:-1], p.stem, p.suffix.lower(), root.name)
        if score <= 0:
            continue
        results.append({"path": str(p), "rel": str(rel).replace(os.sep, "/"), "folder": "/".join(rel.parts[:-1]),
                        "name": p.stem, "ext": p.suffix.lower(), "root": str(root), "score": score})
    results.sort(key=lambda r: (-r["score"], len(r["rel"]), r["rel"]))
    return results


def _score(q: set, fq: str, folder_parts, stem: str, ext: str, root_name: str) -> float:
    """Mức khớp giữa câu nói (tập từ q, câu đã bỏ dấu fq) và một tệp: thư mục chứa, tên tệp, thư mục gốc, định dạng."""
    score = 2.0 * len(q & set(_tokens(" ".join(folder_parts)))) + 1.5 * len(q & set(_tokens(stem)))
    if score <= 0:
        return 0.0
    if folder_parts and fold(folder_parts[-1]) in fq:
        score += 3.0              # tên thư mục xuất hiện nguyên vẹn trong câu nói
    if len(fold(stem)) > 3 and fold(stem) in fq:
        score += 3.0
    root_name = (root_name or "").lower()
    if root_name in _ROOT_WORDS and q & _ROOT_WORDS[root_name]:
        score += 1.0
    if q & _EXT_WORDS.get(ext, set()):
        score += 0.8
    return score


def rank_entries(query: str, entries: List[Dict[str, Any]], exclude: Optional[str] = None) -> List[Dict[str, Any]]:
    """Như find_files nhưng trên danh sách tệp do trình duyệt gửi lên (bản web: thư mục trên máy người dùng).
    entries: [{fid, rel ("thư mục con/tên.pdf"), root (tên thư mục đã chọn), name, ext}]; exclude: fid bỏ qua."""
    q, fq = set(_tokens(query)), fold(query)
    out = []
    for e in entries:
        if exclude and e.get("fid") == exclude:
            continue
        parts = [x for x in str(e.get("rel") or "").split("/") if x][:-1]
        score = _score(q, fq, parts, str(e.get("name") or ""), str(e.get("ext") or "").lower(), str(e.get("root") or ""))
        if score > 0:
            out.append({**e, "folder": "/".join(parts), "score": score})
    out.sort(key=lambda r: (-r["score"], len(r.get("rel") or ""), r.get("rel") or ""))
    return out


# ------------------------------------------------------- tệp gửi lên từ trình duyệt ---
UPLOAD_DIR = Path(os.getenv("DOC_UPLOAD_DIR") or (ROOT / "data" / "uploads"))
UPLOAD_MAX_MB = float(os.getenv("DOC_UPLOAD_MAX_MB", "50") or 50)


def safe_name(name: str) -> str:
    """Tên tệp an toàn để lưu (giữ tiếng Việt, bỏ thư mục và ký tự cấm), giữ đuôi tệp."""
    base = Path(str(name or "").replace("\\", "/")).name
    base = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", base).strip(" .") or "tai-lieu"
    stem, ext = os.path.splitext(base)
    return stem[:120] + ext.lower()


def save_upload(meeting_id: int, name: str, chunks) -> Path:
    """Lưu tệp người dùng gửi lên (để mở / đọc kịch bản) vào thư mục tạm riêng của cuộc họp. Đọc xong thì xóa
    (discard_upload): ảnh từng trang đã lưu trong data/deck_assets, server không giữ tệp gốc của người dùng."""
    fname = safe_name(name)
    ext = os.path.splitext(fname)[1]
    if ext not in EXTS:
        raise ValueError(f"Chưa hỗ trợ định dạng {ext or '(không có đuôi)'}: chỉ mở được {', '.join(EXTS)}")
    d = UPLOAD_DIR / f"m{int(meeting_id)}" / hashlib.sha1(os.urandom(16)).hexdigest()[:12]
    d.mkdir(parents=True, exist_ok=True)
    path, size, limit = d / fname, 0, UPLOAD_MAX_MB * 1024 * 1024
    with open(path, "wb") as fh:
        for chunk in chunks:
            size += len(chunk)
            if size > limit:
                fh.close()
                shutil.rmtree(d, ignore_errors=True)
                raise ValueError(f"Tệp lớn hơn {UPLOAD_MAX_MB:.0f} MB")
            fh.write(chunk)
    return path


def discard_upload(path) -> None:
    p = Path(path)
    if UPLOAD_DIR in p.parents:
        shutil.rmtree(p.parent, ignore_errors=True)


def delete_uploads(meeting_id: int) -> None:
    shutil.rmtree(UPLOAD_DIR / f"m{int(meeting_id)}", ignore_errors=True)


def folders(base: Optional[Path] = None) -> List[str]:
    if base is not None:
        base = Path(base)
        seen = []
        for p in list_files(base):
            f = "/".join(p.relative_to(base).parts[:-1]) or "(gốc)"
            if f not in seen:
                seen.append(f)
        return seen
    return [r.name for r in library_roots()]


def browse(query: str = "", limit: int = 50) -> List[Dict[str, Any]]:
    """Danh sách tệp cho hộp chọn tài liệu: khớp câu tìm (nếu có), mới sửa trước."""
    roots = library_roots()
    if query.strip():
        items = find_files(query)[:limit]
    else:
        items = []
        for p in list_files():
            root = _root_of(p, roots) or p.parent
            rel = p.relative_to(root)
            items.append({"path": str(p), "rel": str(rel).replace(os.sep, "/"), "folder": "/".join(rel.parts[:-1]),
                          "name": p.stem, "ext": p.suffix.lower(), "root": str(root), "score": 0.0})
    for it in items:
        try:
            st = Path(it["path"]).stat()
            it["size"], it["mtime"] = st.st_size, st.st_mtime
        except OSError:
            it["size"], it["mtime"] = 0, 0
        it["root_name"] = Path(it["root"]).name
    if not query.strip():
        items.sort(key=lambda it: -it["mtime"])
    return items[:limit]


# ------------------------------------------------------------ ảnh của slide ---
def assets_key(path: str) -> str:
    """Mã thư mục ảnh của một tệp, theo NỘI DUNG tệp: sửa tệp thì đổi mã; mở lại đúng tệp đó (kể cả gửi lên lại từ
    trình duyệt, nằm ở thư mục tạm khác) thì dùng lại ảnh đã dựng."""
    h = hashlib.sha1(f"v{RENDER_VERSION}|{Path(path).suffix.lower()}|".encode("utf-8"))
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def asset_path(key: str, name: str) -> Optional[Path]:
    """Đường dẫn ảnh đã dựng (chỉ nhận mã và tên hợp lệ, không đi ra ngoài thư mục ảnh)."""
    if not re.fullmatch(r"[0-9a-f]{16}", key or "") or not re.fullmatch(r"\d{1,3}m?\.(png|jpg|jpeg|webp|gif)", name or ""):
        return None
    p = ASSETS_DIR / key / name
    return p if p.is_file() else None


def _asset_url(key: str, name: str) -> str:
    return f"/api/deck-assets/{key}/{name}"


def _first_line(text: str, limit: int = 90) -> str:
    for ln in (text or "").splitlines():
        ln = re.sub(r"\s+", " ", ln).strip()
        if len(ln) >= 3:
            return ln[:limit]
    return ""


def _text_lines(text: str, skip: str = "", n: int = 7) -> List[str]:
    out = []
    for ln in (text or "").splitlines():
        ln = re.sub(r"\s+", " ", ln).strip(" •-\t")
        if len(ln) >= 3 and ln != skip:
            out.append(ln[:220])
        if len(out) >= n:
            break
    return out


def _pdf_slides(path: Path, key: Optional[str]) -> List[Dict[str, Any]]:
    """Mỗi trang PDF thành một slide ảnh (hình trang gốc) kèm chữ của trang."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(path))
    try:
        out_dir = ASSETS_DIR / key if key else None
        if out_dir:
            out_dir.mkdir(parents=True, exist_ok=True)
        slides = []
        for i in range(min(len(pdf), MAX_PAGES)):
            page = pdf[i]
            try:
                tp = page.get_textpage()
                text = tp.get_text_range() or ""
                tp.close()
            except Exception:
                text = ""
            title = _first_line(text) or f"Trang {i + 1}"
            sl: Dict[str, Any] = {"title": title, "layout": "image", "bullets": _text_lines(text, skip=title), "notes": ""}
            if out_dir:
                name = f"{i + 1}.jpg"
                target = out_dir / name
                if not target.exists():
                    w, h = page.get_size()
                    scale = min(3.0, PAGE_PX / max(1.0, w, h))  # cạnh dài khoảng 2000 điểm ảnh: chữ nhỏ vẫn nét
                    img = page.render(scale=scale).to_pil().convert("RGB")
                    # ghi xong mới đổi tên: không để lại ảnh dở khi hai người cùng mở một tệp
                    tmp = out_dir / f".{name}.{os.getpid()}.{threading.get_ident()}.part"
                    img.save(tmp, "JPEG", quality=90, subsampling=0, optimize=True)
                    os.replace(tmp, target)
                sl["image"] = _asset_url(key, name)
            slides.append(sl)
            page.close()
        return slides
    finally:
        pdf.close()


def _soffice() -> Optional[str]:
    """LibreOffice trên máy: SOFFICE_PATH (bản giải nén / cài chỗ khác), PATH, rồi thư mục cài mặc định trên Windows."""
    for c in (os.getenv("SOFFICE_PATH"), shutil.which("soffice"), shutil.which("libreoffice"),
              r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if c and Path(c).is_file():
            return c
    return None


def _run(cmd: List[str], timeout: float) -> subprocess.CompletedProcess:
    """Chạy LibreOffice; quá giờ thì dừng cả cây tiến trình (soffice sinh thêm soffice.bin, chỉ dừng cha là sót con)."""
    kw: Dict[str, Any] = ({"start_new_session": True} if os.name != "nt"
                          else {"creationflags": getattr(subprocess, "CREATE_NO_WINDOW", 0)})
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, **kw)
    try:
        out, err = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        if os.name != "nt":
            try:
                os.killpg(proc.pid, signal.SIGKILL)
            except OSError:
                pass
        else:
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)], capture_output=True)
        proc.kill()
        proc.communicate()
        raise
    return subprocess.CompletedProcess(cmd, proc.returncode, out, err)


def _to_pdf(path: Path) -> Optional[Path]:
    """PowerPoint -> PDF bằng LibreOffice (nếu máy có) để lấy đúng hình từng slide. Mỗi lúc một tệp, hồ sơ LibreOffice
    riêng (không vướng LibreOffice người dùng đang mở), ưu tiên thấp để phòng họp đang ghi lời thoại không bị chậm.
    Không chuyển được thì trả về None và ghi lý do vào log."""
    exe = _soffice()
    if not exe:
        return None
    out = Path(tempfile.mkdtemp(prefix="deck-"))
    cmd = [exe, f"-env:UserInstallation={LO_PROFILE.resolve().as_uri()}", "--headless", "--invisible", "--norestore",
           "--nologo", "--nodefault", "--nolockcheck", "--convert-to", "pdf", "--outdir", str(out), str(path)]
    if os.name != "nt" and shutil.which("nice"):
        cmd = ["nice", "-n", "10"] + cmd
    t0 = time.monotonic()
    try:
        with _LO_LOCK:
            r = _run(cmd, LO_TIMEOUT_S)
    except Exception as e:                       # quá giờ, không chạy được
        log.warning("LibreOffice không chuyển được %s: %s", path.name, e)
        shutil.rmtree(out, ignore_errors=True)
        return None
    pdfs = sorted(out.glob("*.pdf"))
    if not pdfs:
        log.warning("LibreOffice không ra PDF cho %s (mã %s): %s", path.name, r.returncode,
                    (r.stderr or b"").decode("utf-8", "replace")[-300:])
        shutil.rmtree(out, ignore_errors=True)
        return None
    log.info("LibreOffice: %s -> PDF trong %.1f giây", path.name, time.monotonic() - t0)
    return pdfs[0]


def _office_pages(path: Path, key: Optional[str], expect: int) -> Optional[List[Dict[str, Any]]]:
    """Ảnh đúng thiết kế của từng slide PowerPoint (qua LibreOffice), hoặc ảnh đã dựng lần trước cho đúng nội dung tệp
    này. None: không có LibreOffice, chuyển lỗi, hoặc số trang khác số slide (ghép hình với ghi chú sẽ lệch)."""
    if not key or expect <= 0:
        return None
    names = [f"{i}.jpg" for i in range(1, expect + 1)]
    if all((ASSETS_DIR / key / n).is_file() for n in names):
        return [{"image": _asset_url(key, n)} for n in names]
    pdf = _to_pdf(path)
    if pdf is None:
        return None
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(pdf))
        try:
            pages = min(len(doc), MAX_PAGES)
        finally:
            doc.close()
        if pages != expect:
            log.warning("LibreOffice ra %d trang cho %s nhưng tệp có %d slide: không dùng hình", pages, path.name, expect)
            return None
        return _pdf_slides(pdf, key)
    finally:
        shutil.rmtree(pdf.parent, ignore_errors=True)


def _pptx_slides(path: Path, key: Optional[str]) -> List[Dict[str, Any]]:
    """Chữ, ghi chú người trình bày và hình lớn nhất của từng slide PowerPoint (bỏ slide ẩn, như khi trình chiếu và như
    LibreOffice xuất PDF). "_pure": slide chỉ là một hình phủ kín, không chữ (bộ slide xuất từ Canva dạng ảnh)."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(path))
    out_dir = ASSETS_DIR / key if key else None
    slide_area = int(prs.slide_width or 0) * int(prs.slide_height or 0) or 1
    slides = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            else:
                yield sh

    shown = [s for s in prs.slides if s._element.get("show") not in ("0", "false")]
    for i, s in enumerate(shown[:MAX_PAGES]):
        title_shape = s.shapes.title
        title = (title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else "")
        lines, best, best_area = [], None, 0
        for sh in walk(s.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        lines.append(t[:220])
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells)[:220])
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    ext = sh.image.ext.lower()
                    area = int(sh.width or 0) * int(sh.height or 0)
                    if ext in ("png", "jpg", "jpeg", "gif", "webp") and area > best_area:
                        best, best_area = sh, area
                except Exception:
                    pass
        notes = ""
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
            notes = s.notes_slide.notes_text_frame.text.strip()
        uncropped = best is not None and not any(abs(float(getattr(best, c, 0) or 0)) > 1e-3
                                                 for c in ("crop_left", "crop_right", "crop_top", "crop_bottom"))
        pure = uncropped and slide_area > 1 and best_area >= PURE_IMAGE * slide_area and not title and not lines
        if not title and lines:
            title = lines.pop(0)
        sl: Dict[str, Any] = {"title": title or f"Slide {i + 1}", "layout": "bullets", "bullets": lines[:7], "notes": notes,
                              "_pure": pure}
        if best is not None and out_dir:
            out_dir.mkdir(parents=True, exist_ok=True)
            name = f"{i + 1}m.{'jpg' if best.image.ext.lower() == 'jpeg' else best.image.ext.lower()}"   # m: hình trong slide
            target = out_dir / name
            if not target.exists():
                target.write_bytes(best.image.blob)
            # hình phủ gần kín slide (slide xuất từ Canva / Google Slides dạng ảnh): hiện cả hình như trang PDF,
            # chữ và ghi chú vẫn dùng để tìm kiếm, nhắc bài và trình bày
            full = best_area >= FULL_BLEED * slide_area
            sl.update(layout="image" if full else "media", image=_asset_url(key, name))
        slides.append(sl)
    return slides


def _docx_markdown(path: Path) -> str:
    """Word -> markdown: Heading 1 là tên tài liệu, Heading 2 (hoặc Heading 1 tiếp theo) là slide."""
    import docx
    d = docx.Document(str(path))
    out, seen_h1 = [], False
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("title") or (style.startswith("heading 1") and not seen_h1):
            out.append(f"# {t}")
            seen_h1 = True
        elif style.startswith("heading"):
            out.append(f"## {t}")
        elif "list" in style:
            out.append(f"- {t}")
        else:
            out.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append("- " + " | ".join(cells))
    return "\n".join(out)


# ------------------------------------------------------------ đọc tệp ---
def _deck_from_markdown(text: str, fallback_title: str) -> Dict[str, Any]:
    title, slides, cur = fallback_title, [], None
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level, heading = len(m.group(1)), m.group(2).strip()
            if level == 1 and not slides and cur is None:
                title = heading
                continue
            cur = {"title": heading, "bullets": [], "notes": ""}
            slides.append(cur)
            continue
        if cur is None:
            cur = {"title": title, "bullets": [], "notes": ""}
            slides.append(cur)
        b = re.match(r"^\s*(?:[-*+]|\d+[.)])\s+(.*)$", line)
        if b:
            cur["bullets"].append(b.group(1).strip())
        elif line.lstrip().startswith(">"):
            cur["notes"] = (cur["notes"] + " " + line.lstrip("> ").strip()).strip()
        else:
            cur["notes"] = (cur["notes"] + " " + line.strip()).strip()
    # tài liệu không có tiêu đề con: chia đoạn văn thành các slide 5 ý
    if len(slides) == 1 and not slides[0]["bullets"] and len(slides[0]["notes"]) > 600:
        sents = re.split(r"(?<=[.!?])\s+", slides[0]["notes"])
        slides = [{"title": f"{title} ({k // 5 + 1})", "bullets": sents[k:k + 5], "notes": ""}
                  for k in range(0, min(len(sents), 5 * 15), 5)]
    return {"title": title, "slides": slides}


def read_deck(path: str, with_images: bool = True) -> Dict[str, Any]:
    """Đọc một tệp thành dict deck thô {"title", "slides": [...], "source": {...}} (chưa chuẩn hóa).

    with_images=False: không dựng ảnh (chỉ lấy chữ), dùng cho test và xem nhanh."""
    p = Path(path)
    ext = p.suffix.lower()
    source = {"type": "file", "path": str(p), "name": p.name, "ext": ext}
    key = assets_key(str(p)) if with_images and ext in (".pdf", ".pptx", ".ppt") else None
    if ext == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(data, list):
            data = {"title": p.stem, "slides": data}
        if isinstance(data, dict) and not data.get("title"):
            data["title"] = p.stem
    elif ext in (".md", ".txt"):
        data = _deck_from_markdown(p.read_text(encoding="utf-8", errors="replace"), p.stem)
    elif ext == ".docx":
        data = _deck_from_markdown(_docx_markdown(p), p.stem)
    elif ext == ".pdf":
        data = {"title": p.stem, "slides": _pdf_slides(p, key)}
    elif ext == ".pptx":
        slides = _pptx_slides(p, None)            # chữ + ghi chú trước, chưa ghi ảnh
        exact = bool(slides) and all(sl["_pure"] for sl in slides)   # toàn ảnh phủ kín: ảnh gốc là đúng thiết kế
        pages = _office_pages(p, key, len(slides)) if with_images and not exact else None
        if pages:                                 # LibreOffice: đúng hình từng slide, giữ chữ và ghi chú
            for sl, pg in zip(slides, pages):
                sl.update(layout="image", image=pg["image"])
            exact = True
        elif with_images:                         # ảnh gốc của bộ toàn ảnh, hoặc bố cục của ứng dụng + hình lớn nhất
            slides = _pptx_slides(p, key)
        for sl in slides:
            sl.pop("_pure", None)
        source["exact"] = exact and with_images
        data = {"title": p.stem, "slides": slides}
    elif ext == ".ppt":
        pdf = _to_pdf(p)
        if pdf is None:
            raise ValueError("Tệp .ppt cũ cần LibreOffice để mở; anh chị lưu lại thành .pptx hoặc .pdf giúp em")
        data = {"title": p.stem, "slides": _pdf_slides(pdf, key)}
        shutil.rmtree(pdf.parent, ignore_errors=True)
    else:
        raise ValueError(f"Chưa hỗ trợ định dạng {ext}")
    if isinstance(data, dict):
        data["source"] = source
    return data


def read_text(path: str) -> str:
    """Toàn bộ chữ của một tệp (dùng làm kịch bản trình bày): txt, md, docx, pdf, pptx (chữ + ghi chú)."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="replace")
    if ext == ".docx":
        import docx
        return "\n\n".join(par.text.strip() for par in docx.Document(str(p)).paragraphs if par.text.strip())
    if ext == ".pdf":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(p))
        try:
            parts = []
            for i in range(min(len(pdf), 200)):
                tp = pdf[i].get_textpage()
                parts.append(tp.get_text_range() or "")
                tp.close()
            return "\n\n".join(parts)
        finally:
            pdf.close()
    if ext == ".pptx":
        out = []
        for i, sl in enumerate(_pptx_slides(p, None), 1):
            out.append(f"Slide {i}: " + (sl.get("notes") or " ".join([sl["title"]] + sl["bullets"])))
        return "\n\n".join(out)
    if ext == ".json":
        return p.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Chưa đọc được kịch bản dạng {ext}")


# ------------------------------------------------------------ kịch bản theo slide ---
_SCRIPT_MARK = re.compile(r"^\s*[*#>\-]*\s*(?:slide|trang|page|phần|part)\s*(?:số\s*)?(\d{1,3})\s*[*]*\s*[:.\-)\]]*\s*",
                          re.I | re.M)


def split_script(text: str, n_slides: int) -> Optional[List[str]]:
    """Chia kịch bản thành lời cho từng slide, giữ nguyên câu chữ.

    1. Có đánh dấu "Slide 1:", "Trang 2." ... thì chia theo số. 2. Số đoạn (cách nhau dòng trống) đúng bằng số slide thì
    mỗi đoạn một slide. Không chia được thì trả về None (để AI ghép đoạn với slide)."""
    text = (text or "").strip()
    if not text or n_slides <= 0:
        return None
    marks = list(_SCRIPT_MARK.finditer(text))
    if len(marks) >= 2:
        out = [""] * n_slides
        intro = text[:marks[0].start()].strip()
        for k, m in enumerate(marks):
            idx = int(m.group(1)) - 1
            end = marks[k + 1].start() if k + 1 < len(marks) else len(text)
            seg = text[m.end():end].strip()
            if 0 <= idx < n_slides and seg:
                out[idx] = (out[idx] + "\n" + seg).strip()
        if intro and out:
            out[0] = (intro + "\n" + out[0]).strip()
        return out if any(out) else None
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(paras) == n_slides:
        return paras
    if n_slides == 1:
        return [text]
    return None
