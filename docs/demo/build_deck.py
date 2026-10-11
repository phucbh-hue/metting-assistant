"""Bộ slide demo "UrBox Meeting Copilot": HTML -> ảnh (Chromium) -> PPTX và PDF.

- PPTX (slides/demo/): mỗi slide một ảnh phủ kín + ghi chú là lời Jarvis trình bày + chữ nằm dưới ảnh để ứng dụng tìm
  kiếm và nhắc bài. Ứng dụng hiện cả hình (kể cả máy không có LibreOffice) và đọc ghi chú khi chọn "Theo ghi chú
  trong tệp".
- PDF (docs/demo/): cùng nội dung, chữ chọn được, để gửi người tham dự.
- docs/demo/kich-ban-demo.md: mục 5 (thông điệp, câu hỏi dự kiến, lời Jarvis của từng slide) viết lại từ SLIDES.

Phong cách UrBox (lấy từ urbox.vn): nền trắng / #F5F5FA, tím #6821C2 chủ đạo, tím #8235E4, vàng #FFB700 điểm nhấn,
thẻ bo 20px, font Be Vietnam Pro. Bìa và slide demo dùng dải chuyển màu tím như trang chủ UrBox.

Chạy: python docs/demo/build_deck.py   (cần playwright + chromium, python-pptx, Pillow; tải font Be Vietnam Pro từ
Google Fonts). Sửa nội dung: sửa SLIDES (title, body, notes, message, qa, hidden) rồi chạy lại.
"""
import json
import tempfile
from pathlib import Path

from playwright.sync_api import sync_playwright

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
IMG = (REPO / "docs" / "huong-dan" / "img").as_uri()
D = (HERE / "img").as_uri()
MASCOT = (REPO / "static" / "mascot" / "urbox-mascot.svg").as_uri()
SCRIPT_MD = HERE / "kich-ban-demo.md"
MARK = "<!-- LOI_TRINH_BAY -->"
NAME = "UrBox Meeting Copilot - Demo 10.2026"

I = {  # biểu tượng nét (24x24), cùng kiểu biểu tượng nét tím trên urbox.vn
    "mic": '<path d="M12 2a3 3 0 0 0-3 3v7a3 3 0 0 0 6 0V5a3 3 0 0 0-3-3z"/><path d="M19 10v2a7 7 0 0 1-14 0v-2"/><path d="M12 19v3"/>',
    "file": '<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><path d="M14 2v6h6"/><path d="M16 13H8"/><path d="M16 17H8"/>',
    "clock": '<circle cx="12" cy="12" r="10"/><path d="M12 6v6l4 2"/>',
    "history": '<path d="M3 12a9 9 0 1 0 3-6.7L3 8"/><path d="M3 3v5h5"/><path d="M12 7v5l3 2"/>',
    "ask": '<path d="M21 15a2 2 0 0 1-2 2H7l-4 4V5a2 2 0 0 1 2-2h14a2 2 0 0 1 2 2z"/><path d="M9.5 9a2.5 2.5 0 1 1 3.5 2.3c-.6.3-1 .9-1 1.6"/><path d="M12 16h.01"/>',
    "users": '<path d="M17 21v-2a4 4 0 0 0-4-4H5a4 4 0 0 0-4 4v2"/><circle cx="9" cy="7" r="4"/><path d="M23 21v-2a4 4 0 0 0-3-3.87"/><path d="M16 3.13a4 4 0 0 1 0 7.75"/>',
    "user": '<circle cx="12" cy="8" r="4"/><path d="M4 21v-1a6 6 0 0 1 6-6h4a6 6 0 0 1 6 6v1"/>',
    "shield": '<path d="M12 22s8-4 8-10V5l-8-3-8 3v7c0 6 8 10 8 10z"/>',
    "lock": '<rect x="3" y="11" width="18" height="11" rx="2"/><path d="M7 11V7a5 5 0 0 1 10 0v4"/>',
    "cloud": '<path d="M18 10h-1.26A8 8 0 1 0 9 20h9a5 5 0 0 0 0-10z"/>',
    "calendar": '<rect x="3" y="4" width="18" height="18" rx="2"/><path d="M16 2v4M8 2v4M3 10h18"/>',
    "search": '<circle cx="11" cy="11" r="8"/><path d="M21 21l-4.35-4.35"/>',
    "spark": '<path d="M12 3l1.9 5.8L20 10l-6.1 1.2L12 17l-1.9-5.8L4 10l6.1-1.2z"/>',
    "eyeoff": '<path d="M17.94 17.94A10.07 10.07 0 0 1 12 20c-7 0-11-8-11-8a18.45 18.45 0 0 1 5.06-5.94"/><path d="M9.9 4.24A9.12 9.12 0 0 1 12 4c7 0 11 8 11 8a18.5 18.5 0 0 1-2.16 3.19"/><path d="M1 1l22 22"/>',
    "folder": '<path d="M22 19a2 2 0 0 1-2 2H4a2 2 0 0 1-2-2V5a2 2 0 0 1 2-2h5l2 3h9a2 2 0 0 1 2 2z"/>',
    "brief": '<rect x="2" y="7" width="20" height="14" rx="2"/><path d="M16 21V5a2 2 0 0 0-2-2h-4a2 2 0 0 0-2 2v16"/>',
    "check": '<path d="M20 6L9 17l-5-5"/>',
    "video": '<rect x="2" y="6" width="14" height="12" rx="2"/><path d="M16 10l6-4v12l-6-4z"/>',
    "wave": '<path d="M2 12h2M6 8v8M10 5v14M14 8v8M18 10v4M22 12h0"/>',
    "server": '<rect x="2" y="3" width="20" height="8" rx="2"/><rect x="2" y="13" width="20" height="8" rx="2"/><path d="M6 7h.01M6 17h.01"/>',
    "db": '<ellipse cx="12" cy="5" rx="8" ry="3"/><path d="M4 5v14c0 1.7 3.6 3 8 3s8-1.3 8-3V5"/><path d="M4 12c0 1.7 3.6 3 8 3s8-1.3 8-3"/>',
    "plug": '<path d="M9 2v6M15 2v6M6 8h12v3a6 6 0 0 1-12 0z"/><path d="M12 17v5"/>',
    "gauge": '<path d="M12 14l4-4"/><path d="M3.3 19a10 10 0 1 1 17.4 0"/>',
    "coin": '<circle cx="12" cy="12" r="9"/><path d="M14.5 9a2.5 2 0 0 0-2.5-1.5c-1.5 0-2.5.8-2.5 2s1 1.7 2.5 2 2.5.8 2.5 2-1 2-2.5 2A2.5 2 0 0 1 9.5 15"/><path d="M12 6v1.5M12 16.5V18"/>',
    "layers": '<path d="M12 2l10 5-10 5L2 7z"/><path d="M2 17l10 5 10-5"/><path d="M2 12l10 5 10-5"/>',
}


def ic(name, size=30, color="#6821C2", box=64):
    svg = (f'<svg width="{size}" height="{size}" viewBox="0 0 24 24" fill="none" stroke="{color}" stroke-width="2" '
           f'stroke-linecap="round" stroke-linejoin="round">{I[name]}</svg>')
    return f'<span class="ic" style="width:{box}px;height:{box}px">{svg}</span>'


def shot(src, w, h, pos="50% 0%"):
    return f'<div class="shot" style="width:{w}px;height:{h}px"><img src="{src}" style="object-position:{pos}"></div>'


def row(icon, head, text):
    return f'<div class="row">{ic(icon)}<div><div class="rh">{head}</div><div class="rt">{text}</div></div></div>'


K = '<span class="cite">K1</span>'
M = '<span class="cite">M1</span>'

SLIDES = [
    # 1 --------------------------------------------------------------------------------------------------------
    dict(kind="hero", title="UrBox Meeting Copilot", body=f"""
      <div class="cl">
        <span class="pill-gold">Demo nội bộ · 10/2026</span>
        <h1>UrBox Meeting Copilot</h1>
        <p class="sub">Trợ lý AI hỗ trợ ngay trong lúc họp: gợi ý trả lời khách, gợi ý câu hỏi phỏng vấn, biết ai đang nói</p>
        <p class="gold">Bộ slide này do trợ lý Jarvis tự mở và trình bày</p>
        <p class="who">Phụ trách: phuc.bh@urbox.vn</p>
      </div>
      <div class="cr"><div class="halo"></div><img class="mascot" src="{MASCOT}"></div>""",
         hidden=["Trợ lý AI hỗ trợ ngay trong lúc họp", "Gợi ý trả lời khách, gợi ý câu hỏi phỏng vấn, biết ai đang nói"],
         message="Meeting Copilot hỗ trợ ngay trong lúc họp, không chỉ ghi chép sau buổi họp.",
         qa=[("Cái này khác gì Gemini ghi chú trong Google Meet hay Otter, Fireflies?",
              "Các công cụ đó mạnh ở chép lời và tóm tắt sau buổi họp. Meeting Copilot tập trung vào lúc đang họp, dựa "
              "trên tài liệu và lịch sử họp của chính UrBox, và chạy cả khi họp trực tiếp. Slide 3 nói rõ 5 điểm khác.")],
         notes="Dạ em chào anh chị. Em là Jarvis, trợ lý họp của UrBox. Hôm nay em tự trình bày về chính mình: UrBox "
               "Meeting Copilot. Điều em muốn anh chị nhớ nhất là: em hỗ trợ ngay trong lúc đang họp, không chỉ ghi chép "
               "sau buổi họp. Em sẽ đi qua những điểm khác biệt, hai tình huống dùng thực tế là họp với khách và phỏng vấn, "
               "cách em hoạt động về kỹ thuật, rồi đến phần demo trực tiếp."),
    # 2 --------------------------------------------------------------------------------------------------------
    dict(eyebrow="Vấn đề", title="Ghi chú sau buổi họp đã có nhiều công cụ. Lúc cần trợ giúp nhất lại là trong buổi họp",
         body=f"""
      <div class="cols2 problem">
        <div class="card">
          <div class="ch muted">Đã có công cụ làm tốt</div>
          <div class="chips"><span>Gemini trong Google Meet</span><span>Otter</span><span>Fireflies</span><span>tl;dv</span></div>
          <ul class="ticks">
            <li>Chép lời, tóm tắt</li><li>Liệt kê việc cần làm</li><li>Tìm lại nội dung các cuộc họp</li>
          </ul>
          <p class="foot-note">Mạnh ở phần sau buổi họp, chủ yếu với họp online.</p>
        </div>
        <div class="card white">
          <div class="ch">Còn thiếu: trợ giúp ngay trong buổi</div>
          <div class="rows tight">
            {row('brief', 'Khách hỏi điều BD chưa nắm', 'Phải hẹn trả lời sau, hoặc mở tài liệu ngay trước mặt khách.')}
            {row('user', 'Ứng viên nói thuật ngữ lạ', 'Người phỏng vấn cần câu hỏi đào sâu, nhưng không phải chuyên gia mọi mảng.')}
            {row('history', 'Nói khác điều đã chốt buổi trước', 'Không ai nhớ buổi nào đã chốt gì.')}
            {row('users', 'Họp trực tiếp, nhiều người', 'Biên bản không rõ ai nói câu nào, người mới là ai.')}
          </div>
        </div>
      </div>""",
         hidden=["Công cụ ghi chú: chép lời, tóm tắt sau buổi họp", "Thiếu: trợ giúp ngay trong buổi họp",
                 "Khách hỏi điều BD chưa nắm", "Ứng viên nói thuật ngữ lạ", "Nói khác điều đã chốt",
                 "Họp trực tiếp không rõ ai nói"],
         message="Khoảng trống không nằm ở ghi chép, mà ở trợ giúp ngay lúc đang họp, bằng dữ liệu riêng của UrBox.",
         qa=[("Đây có phải vấn đề thật không, hay dùng Gemini là đủ?",
              "Nếu chỉ cần biên bản thì công cụ sẵn có là đủ. Bốn tình huống bên phải xảy ra giữa buổi họp và cần dữ "
              "liệu riêng của UrBox (bảng giá, JD, lịch sử họp, hồ sơ giọng), công cụ bên ngoài không có dữ liệu này.")],
         notes="Ghi chú họp hiện có nhiều công cụ tốt, như Gemini trong Google Meet, Otter hay Fireflies: chép lời, tóm "
               "tắt, liệt kê việc cần làm sau buổi họp. Nhưng lúc cần trợ giúp nhất lại xảy ra ngay giữa buổi: khách hỏi một "
               "điều đội BD chưa nắm, ứng viên nhắc tới một kỹ thuật mà người phỏng vấn muốn hỏi sâu hơn, buổi này nói khác "
               "điều đã chốt ở buổi trước, hay họp trực tiếp đông người mà biên bản không rõ ai nói gì. Những lúc đó cần dữ "
               "liệu riêng của UrBox, và cần ngay."),
    # 3 --------------------------------------------------------------------------------------------------------
    dict(eyebrow="Khác biệt", title="5 điểm khác so với các công cụ ghi chú họp phổ biến", body=f"""
      <div class="diff">
        <div class="drow top"><span class="dn">1</span><div><div class="dh">Gợi ý trả lời ngay khi khách hỏi <span class="badge-gold">Ưu tiên số 1</span></div>
          <div class="dt">Lấy từ tài liệu và các buổi họp trước của UrBox, có dấu nguồn, trên bảng riêng chỉ đội mình thấy</div></div></div>
        <div class="drow"><span class="dn">2</span><div><div class="dh">Trợ lý riêng khi phỏng vấn</div>
          <div class="dt">Gợi ý câu hỏi đào sâu theo JD, giải thích thuật ngữ, tra thêm trên mạng có nguồn</div></div></div>
        <div class="drow"><span class="dn">3</span><div><div class="dh">Nhận ra người nói bằng mẫu giọng</div>
          <div class="dt">Cả khi họp trực tiếp; người mới được gắn tên qua hội thoại; hồ sơ giọng dùng chung toàn công ty</div></div></div>
        <div class="drow"><span class="dn">4</span><div><div class="dh">Nhớ xuyên các cuộc họp</div>
          <div class="dt">Biên bản tự nhắc khi buổi này nói khác điều đã chốt ở buổi trước</div></div></div>
        <div class="drow"><span class="dn">5</span><div><div class="dh">Tiếng Việt, dữ liệu trên hạ tầng UrBox</div>
          <div class="dt">Gọi trợ lý bằng giọng nói tiếng Việt; dữ liệu ở MongoDB Atlas của UrBox, đăng nhập @urbox.vn</div></div></div>
      </div>
      <p class="bottom-note">Chép lời, biên bản, bot Google Meet, lưu Google Drive là phần cơ bản, công cụ khác cũng có (slide 12).</p>""",
         hidden=["Gợi ý trả lời ngay khi khách hỏi, từ tài liệu và buổi họp trước", "Trợ lý riêng khi phỏng vấn",
                 "Nhận ra người nói bằng mẫu giọng", "Nhớ xuyên các cuộc họp", "Tiếng Việt, dữ liệu trên hạ tầng UrBox"],
         message="Khác biệt nằm ở hỗ trợ ngay lúc họp, dựa trên dữ liệu riêng của UrBox; phần ghi chép là cơ bản.",
         qa=[("Trong 5 điểm này, điểm nào công cụ khác thật sự không làm được?",
              "Điểm 1, 2, 4 cần tài liệu nội bộ và lịch sử họp của UrBox, công cụ bên ngoài không có. Điểm 3 dùng hồ sơ "
              "giọng chung toàn công ty, chạy cả họp trực tiếp không cần bot. Mình không khẳng định không công cụ nào có "
              "từng phần; điểm mạnh là gộp cả 5 trên dữ liệu của mình.")],
         notes="Em có năm điểm khác so với các công cụ ghi chú họp phổ biến. Một, khi khách hỏi, em gợi ý câu trả lời ngay, "
               "lấy từ tài liệu và các buổi họp trước của UrBox, có dấu nguồn, trên một bảng riêng chỉ đội mình thấy. Đây là "
               "tính năng em ưu tiên nhất. Hai, khi phỏng vấn, em là trợ lý riêng của người phỏng vấn. Ba, em nhận ra người "
               "nói bằng mẫu giọng, kể cả khi họp trực tiếp. Bốn, em nhớ xuyên các cuộc họp. Năm, em nói tiếng Việt và dữ "
               "liệu nằm trên hạ tầng của UrBox. Còn chép lời, biên bản là phần cơ bản, em để ở cuối."),
    # 4 --------------------------------------------------------------------------------------------------------
    dict(eyebrow="Use case 1 · Họp với khách", title="Khách hỏi điều BD chưa nắm: vài giây sau đã có gợi ý trả lời, kèm nguồn",
         body=f"""
      <div class="split">
        <div class="flowv">
          <div class="step"><div class="sl">Khách hỏi</div>
            <div class="bubble">"Phí tích hợp API bao nhiêu, có khác lần trước không?"</div></div>
          <div class="step"><div class="sl">Sau 4-8 giây <span class="badge-gold sm">đo thật</span></div>
            <div class="answer">5.000.000đ {K}, thu một lần, gồm hỗ trợ 24/7
              <div class="note">Ghi chú: buổi 02/10 đội mình đã báo 4.000.000đ {M}, nên giải thích phần chênh lệch</div></div></div>
          <div class="step"><div class="sl">BD trả lời ngay</div>
            <div class="plain">Không phải hẹn khách "em kiểm tra rồi gửi lại", và biết thông tin lấy từ đâu.</div></div>
        </div>
        {shot(D + '/bd-panel.png', 1000, 597)}
      </div>""",
         hidden=["Khách hỏi: phí tích hợp API bao nhiêu", "Sau 4-8 giây: gợi ý 5.000.000đ, nguồn bảng giá",
                 "Ghi chú: buổi trước đã báo 4.000.000đ", "BD trả lời ngay, có nguồn"],
         message="BD trả lời được ngay cả câu chưa nắm, và biết nguồn để tin.",
         qa=[("AI có bịa giá hay chính sách không?",
              "Chỉ dùng nguồn trong tài liệu nhóm và các buổi trước; mỗi ý có dấu K (tài liệu) hoặc M (cuộc họp trước), "
              "bấm xem được. Thiếu thông tin thì gợi ý 'bên em kiểm tra và gửi lại', độ chắc thấp."),
             ("Khách có nhìn thấy không?",
              "Không. Bảng mở trên máy đội BD, không lên màn hình chung, trợ lý không đọc thành tiếng.")],
         notes="Tình huống đầu tiên, cũng là tính năng em ưu tiên nhất: họp với khách. Khách hỏi: phí tích hợp API bao nhiêu, "
               "có khác lần trước không? Chỉ vài giây sau, trên bảng riêng của đội BD đã có gợi ý: năm triệu đồng, thu một "
               "lần, lấy từ bảng giá. Em còn ghi chú: buổi ngày hai tháng mười, đội mình đã báo bốn triệu, nên giải thích "
               "phần chênh lệch. Đội BD trả lời được ngay, kể cả khi chưa nắm câu đó, và biết thông tin lấy từ đâu."),
    # 5 --------------------------------------------------------------------------------------------------------
    dict(eyebrow="Use case 1 · Cách hoạt động", title="Câu trả lời chỉ lấy từ nguồn của UrBox, và chỉ đội mình thấy", body=f"""
      <div class="flowh">
        <div class="fcard"><span class="fn">1</span>{ic('ask')}<div class="fh">Nhận ra câu hỏi</div>
          <div class="ft">Chỉ câu của người đánh dấu là Khách; bỏ lời chào, câu xác nhận</div></div>
        <div class="fcard"><span class="fn">2</span>{ic('search')}<div class="fh">Tìm nguồn</div>
          <div class="ft">Tài liệu nhóm (bảng giá, hồ sơ năng lực, FAQ) và biên bản, lời nói các buổi trước</div></div>
        <div class="fcard"><span class="fn">3</span>{ic('spark')}<div class="fh">Viết gợi ý</div>
          <div class="ft">Câu trả lời ngắn, dấu nguồn K / M, ghi chú cho đội, độ chắc</div></div>
        <div class="fcard"><span class="fn">4</span>{ic('cloud')}<div class="fh">Tra thêm trên mạng</div>
          <div class="ft">Chỉ khi bấm; ghi rõ "chưa kiểm chứng", nguồn W bấm mở trang</div></div>
      </div>
      <div class="cols2 lower">
        <div class="card">{ic('users')}<div><div class="rh">Chủ nhóm hỏi trên toàn bộ cuộc họp</div>
          <div class="rt">"Báo cáo các câu phàn nàn của khách", "Mình đã hứa gì với khách?"</div></div></div>
        <div class="card">{ic('coin')}<div><div class="rh">Khoảng 250đ mỗi câu gợi ý</div>
          <div class="rt">Đo thật: 0,007-0,013 USD, 4-8 giây; tìm kiếm chạy ngay trên server</div></div></div>
      </div>""",
         hidden=["Nhận ra câu hỏi của khách", "Tìm trong tài liệu nhóm và các buổi trước", "Viết gợi ý có nguồn",
                 "Tra thêm trên mạng khi bấm", "Chủ nhóm hỏi trên toàn bộ cuộc họp", "Khoảng 250đ mỗi câu gợi ý"],
         message="Gợi ý có căn cứ, có kiểm soát, chi phí thấp.",
         qa=[("Tìm kiếm làm thế nào, có cần vector database không?",
              "Không. Tìm theo từ khóa (BM25, bỏ dấu tiếng Việt, ghép cặp từ) chạy ngay trên server, không gửi tài liệu đi "
              "lập chỉ mục ở dịch vụ ngoài. Đổi sang tìm theo nghĩa khi cần là việc có thể làm sau."),
             ("Mỗi câu tốn bao nhiêu, có gọi AI liên tục không?",
              "Đo 4 lượt với Claude Sonnet 5.5: 0,007-0,013 USD mỗi câu, 4-8 giây. Chỉ gọi AI khi câu có dấu hiệu là câu "
              "hỏi, tối đa một câu tự nhận ra mỗi 6 giây.")],
         notes="Về cách hoạt động: em chỉ xét câu của người được đánh dấu là khách, bỏ qua lời chào và câu xác nhận. Em tìm "
               "trong tài liệu của nhóm như bảng giá, hồ sơ năng lực, và trong biên bản, lời nói của các buổi trước. Rồi em "
               "viết gợi ý ngắn, có dấu nguồn, kèm ghi chú cho đội và mức độ chắc chắn. Cần thông tin bên ngoài thì đội bấm "
               "tra thêm trên mạng. Việc tìm kiếm chạy ngay trên server, không cần cơ sở dữ liệu vector. Mỗi câu gợi ý tốn "
               "khoảng hai trăm năm mươi đồng."),
    # 6 --------------------------------------------------------------------------------------------------------
    dict(eyebrow="Use case 2 · Phỏng vấn", title="Phỏng vấn: gợi ý câu hỏi đào sâu và giải thích thuật ngữ ngay trong buổi",
         body=f"""
      <div class="split">
        <div class="leftcol">
          {shot(D + '/interview-panel.png', 960, 576)}
          <p class="cap">Cách dùng: nhóm tuyển dụng bật chế độ BD, tải JD và tiêu chí đánh giá, đánh dấu ứng viên là Khách.</p>
        </div>
        <div class="rows">
          {row('ask', 'Gợi ý câu hỏi đào sâu', '"Gợi ý 3 câu hỏi khai thác sâu về sự cố ứng viên vừa kể": 3 câu bám tiêu chí trong JD.')}
          {row('search', 'Thông tin ngành tức thì', '"Outbox pattern là gì, trả lời vậy có hợp lý không?": giải thích, tra thêm trên mạng có nguồn.')}
          {row('user', 'Trả lời câu hỏi của ứng viên', 'Ứng viên hỏi chế độ làm việc: trả lời theo JD, có nguồn.')}
          <div class="callout">Thử với AI thật: 4 câu chạy song song, xong trong <span class="nw">7-9 giây</span></div>
        </div>
      </div>""",
         hidden=["Gợi ý câu hỏi đào sâu theo JD", "Giải thích thuật ngữ, tra thêm trên mạng",
                 "Trả lời câu hỏi của ứng viên theo JD", "4 câu song song trong 7-9 giây"],
         message="Người phỏng vấn có trợ lý riêng, hỏi sâu hơn mà không cần là chuyên gia mọi mảng.",
         qa=[("Team Engineering tuyển người có dùng được ngay không?",
              "Được, với tính năng hiện có: tạo nhóm tuyển dụng, bật chế độ BD, tải JD và tiêu chí. Chưa có giao diện "
              "riêng cho phỏng vấn: nhãn trên bảng vẫn là Khách / Đội mình."),
             ("Ứng viên có biết không, có vấn đề dữ liệu cá nhân không?",
              "Bảng mở trên máy người phỏng vấn, không đọc thành tiếng. Ghi âm và lưu mẫu giọng chỉ khi ứng viên đồng ý, "
              "theo Nghị định 13/2023.")],
         notes="Tình huống thứ hai là phỏng vấn. Ứng viên kể đã xử lý sự cố mất message bằng outbox pattern. Người phỏng vấn "
               "gõ vào bảng riêng: gợi ý ba câu hỏi khai thác sâu. Em trả lời ngay ba câu bám theo tiêu chí trong mô tả công "
               "việc, như cách phát hiện sự cố và cách thiết kế chống ghi trùng. Gặp thuật ngữ chưa rõ, người phỏng vấn hỏi "
               "em, em giải thích và tra thêm trên mạng có nguồn. Khi ứng viên hỏi về chế độ làm việc, em trả lời theo mô tả "
               "công việc. Tính năng này dùng được ngay cho các buổi tuyển dụng của team Engineering."),
    # 7 --------------------------------------------------------------------------------------------------------
    dict(eyebrow="Nhận ra người nói", title="Biết ai đang nói nhờ mẫu giọng, kể cả người lần đầu tham gia", body=f"""
      <div class="vflow">
        <div class="vnode">{ic('wave')}<div class="vh">Mỗi câu nói</div><div class="vt">Soniox tách người nói trong phiên</div></div>
        <div class="varrow">›</div>
        <div class="vnode">{ic('layers')}<div class="vh">Vector giọng 192 số</div><div class="vt">ERes2NetV2, chạy trên CPU server</div></div>
        <div class="varrow">›</div>
        <div class="vnode">{ic('users')}<div class="vh">So với hồ sơ giọng</div><div class="vt">Dùng chung toàn công ty</div></div>
        <div class="varrow">›</div>
        <div class="vnode hl">{ic('check')}<div class="vh">Khớp: hiện đúng tên</div><div class="vt">Kèm số câu, thời lượng nói</div></div>
      </div>
      <div class="vbottom">
        <div class="pair">
          <figure>{shot(D + '/voice-before.png', 200, 361)}<figcaption>Lần đầu</figcaption></figure>
          <span class="parrow">›</span>
          <figure>{shot(D + '/voice-after.png', 200, 361)}<figcaption>Đã có mẫu giọng</figcaption></figure>
        </div>
        <div class="card newp">
          <div class="ch">Người mới</div>
          <ol class="steps">
            <li>Hiện là "Người nói N"</li>
            <li>Tự giới thiệu hoặc được gọi tên: AI đoán tên, chắc từ 70% thì tự gắn, thấp hơn thì hỏi lại</li>
            <li>Người đó đồng ý: lưu mẫu giọng, buổi sau tự nhận ra</li>
          </ol>
        </div>
        <div class="card stat">
          <div class="big">97% · 98%</div>
          <div class="rt">câu gán đúng người, đo trên 2 cuộc họp thật nội bộ (4 người và 2 người)</div>
          <div class="mini">{ic('lock', 22, box=40)}<span>Chỉ lưu vector, không lưu âm thanh mẫu</span></div>
        </div>
      </div>""",
         hidden=["Mỗi câu nói tính vector giọng 192 số", "So với hồ sơ giọng nói toàn công ty", "Khớp thì hiện đúng tên",
                 "Người mới: AI đoán tên, lưu mẫu giọng khi đồng ý", "97% và 98% câu gán đúng người",
                 "Chỉ lưu vector, không lưu âm thanh"],
         message="Biên bản biết chính xác ai nói gì; người mới cũng được nhận ra từ buổi sau.",
         qa=[("Độ chính xác bao nhiêu, đo thế nào?",
              "Đo trên 2 cuộc họp thật nội bộ đã gán nhãn tay: 97% và 98% câu gán đúng người. Dữ liệu còn ít, cần đo thêm "
              "trên nhiều cuộc họp và nhiều loại mic."),
             ("Người mới thì xử lý thế nào?",
              "Lần đầu hiện là Người nói N. Khi họ tự giới thiệu hoặc được gọi tên, AI đoán tên: chắc từ 70% thì tự gắn, "
              "thấp hơn thì hiện gợi ý để xác nhận. Người đó đồng ý thì lưu mẫu giọng, buổi sau tự nhận ra."),
             ("Dữ liệu giọng lưu thế nào, có xóa được không?",
              "Mỗi người một vector 192 số, không lưu âm thanh. Chỉ lưu khi có đồng ý theo Nghị định 13/2023; xóa hồ sơ bất "
              "cứ lúc nào.")],
         notes="Phần nhận ra người nói. Với mỗi câu nói, em tính một vector giọng gồm một trăm chín mươi hai con số, chạy ngay "
               "trên server, rồi so với hồ sơ giọng nói của công ty. Khớp thì em hiện đúng tên, kèm số câu và thời lượng nói "
               "của từng người. Người lần đầu tham gia sẽ hiện là người nói mới; khi họ tự giới thiệu hoặc được gọi tên, em "
               "đoán tên và gắn vào, chưa chắc thì em hỏi lại. Nếu người đó đồng ý, em lưu mẫu giọng, từ buổi sau em tự nhận "
               "ra. Đo trên hai cuộc họp thật nội bộ, em gán đúng người nói chín mươi bảy và chín mươi tám phần trăm số câu."),
    # 8 --------------------------------------------------------------------------------------------------------
    dict(eyebrow="Thời gian thực", title="Mọi thứ chạy trong lúc họp: số đo thực tế", body=f"""
      <div class="stats">
        <div class="tile">{ic('mic')}<div class="v">Ngay khi nói</div><div class="l">Chữ hiện theo thời gian thực (Soniox, tiếng Việt)</div></div>
        <div class="tile">{ic('layers')}<div class="v">0,11 giây</div><div class="l">Tính vector giọng cho mỗi giây âm thanh, trên CPU</div></div>
        <div class="tile">{ic('ask')}<div class="v">4-8 giây</div><div class="l">Gợi ý trả lời khi khách hỏi</div></div>
        <div class="tile">{ic('gauge')}<div class="v">7-9 giây</div><div class="l">4 câu hỏi cùng lúc, chạy song song</div></div>
        <div class="tile">{ic('cloud')}<div class="v">15-25 giây</div><div class="l">Tra thêm trên mạng, có nguồn</div></div>
        <div class="tile">{ic('coin')}<div class="v">~250đ</div><div class="l">Một câu gợi ý (0,007-0,013 USD)</div></div>
      </div>
      <p class="bottom-note">Đo ngày 10-11/10/2026 với Claude Sonnet 5.5, dữ liệu mẫu, mạng văn phòng. Độ trễ đổi theo mạng và tải của dịch vụ AI.</p>""",
         hidden=["Chữ hiện theo thời gian thực", "Vector giọng 0,11 giây cho mỗi giây âm thanh",
                 "Gợi ý trả lời 4-8 giây", "4 câu song song 7-9 giây", "Tra thêm trên mạng 15-25 giây",
                 "Khoảng 250đ mỗi câu gợi ý"],
         message="Đủ nhanh để dùng ngay trong cuộc nói chuyện, chi phí mỗi lượt thấp.",
         qa=[("Số đo ở đâu ra, lặp lại được không?",
              "Đo bằng script chạy AI thật với dữ liệu mẫu (có trong repo); chạy lại được bất cứ lúc nào. Chưa đo khi "
              "nhiều phòng họp chạy cùng lúc.")],
         notes="Về thời gian thực, em đưa số đo thực tế. Chữ hiện ngay khi đang nói. Tính vector giọng chỉ mất khoảng không "
               "phẩy mười một giây cho mỗi giây âm thanh. Gợi ý trả lời khi khách hỏi mất từ bốn đến tám giây. Bốn câu hỏi cùng "
               "lúc xong trong bảy đến chín giây vì chạy song song. Tra thêm trên mạng mất mười lăm đến hai mươi lăm giây. Mỗi "
               "câu gợi ý tốn khoảng hai trăm năm mươi đồng. Các số này đo với dữ liệu mẫu, sẽ thay đổi theo mạng và tải."),
    # 9 --------------------------------------------------------------------------------------------------------
    dict(eyebrow="Kỹ thuật", title="Kiến trúc gọn: một server, dịch vụ AI qua API, dữ liệu ở Atlas của UrBox", body=f"""
      <div class="arch">
        <div class="acol">
          <div class="abox">{ic('mic', 26, box=52)}<div><b>Mic trình duyệt</b><span>Họp trực tiếp</span></div></div>
          <div class="abox">{ic('video', 26, box=52)}<div><b>Bot Google Meet</b><span>Recall, mỗi người một luồng</span></div></div>
        </div>
        <div class="aarrow"><span>WebSocket</span>›</div>
        <div class="acenter">
          {ic('server', 34, box=76)}
          <div class="atitle">Server Meeting Copilot</div>
          <div class="asub">FastAPI (Python), Docker trên Render</div>
          <ul><li>Tách người nói, vector giọng ERes2NetV2</li><li>Tìm kiếm BM25 tài liệu và lịch sử họp</li>
          <li>Trợ lý AI, bảng BD, biên bản</li></ul>
        </div>
        <div class="aarrow">‹ ›</div>
        <div class="acol right">
          <div class="abox">{ic('wave', 26, box=52)}<div><b>Soniox</b><span>Chép lời, giọng đọc</span></div></div>
          <div class="abox">{ic('spark', 26, box=52)}<div><b>Claude Sonnet 5.5</b><span>API key công ty hoặc gói cá nhân</span></div></div>
          <div class="abox">{ic('db', 26, box=52)}<div><b>MongoDB Atlas</b><span>Cuộc họp, hồ sơ giọng</span></div></div>
          <div class="abox">{ic('folder', 26, box=52)}<div><b>Google Drive, Lịch</b><span>Biên bản, lịch họp</span></div></div>
          <div class="abox">{ic('plug', 26, box=52)}<div><b>Dữ liệu nội bộ qua MCP</b><span>Jira: đang dùng dữ liệu mẫu</span></div></div>
        </div>
      </div>
      <div class="chips bottom"><span>376 test tự động</span><span>Đăng nhập Google @urbox.vn</span><span>Đổi được nhà cung cấp AI: Claude, Gemini, gói ChatGPT</span></div>""",
         hidden=["Mic trình duyệt và bot Google Meet qua WebSocket", "Server FastAPI, Docker trên Render",
                 "Soniox, Claude Sonnet 5.5, MongoDB Atlas, Google Drive", "Dữ liệu nội bộ qua MCP", "376 test tự động"],
         message="Kiến trúc đơn giản, dễ vận hành, không khóa chặt vào một nhà cung cấp.",
         qa=[("Một server có phải điểm nghẽn không?",
              "Phòng họp đang diễn ra giữ trong bộ nhớ nên hiện chạy một bản. Chưa đo tải nhiều phòng cùng lúc; muốn mở "
              "rộng cần tách trạng thái phòng họp ra khỏi bộ nhớ."),
             ("Phụ thuộc những nhà cung cấp nào?",
              "Soniox (chép lời), Anthropic (AI), Recall (bot Meet). Phần AI đã đổi được sang Gemini hoặc gói ChatGPT; "
              "giọng đọc có bản Piper chạy ngay trên server."),
             ("Kiểm thử và triển khai thế nào?",
              "376 test tự động; bản web cập nhật từ nhánh release, chỉ đẩy khi không có cuộc họp nào đang diễn ra.")],
         notes="Phần này dành cho anh chị bên kỹ thuật. Âm thanh đến từ mic trên trình duyệt khi họp trực tiếp, hoặc từ bot "
               "Google Meet với mỗi người một luồng riêng, gửi qua WebSocket tới một server FastAPI chạy Docker trên Render. "
               "Server tự tách người nói, tính vector giọng và tìm kiếm tài liệu. Chép lời dùng Soniox, phần AI dùng Claude "
               "Sonnet năm phẩy năm qua API của công ty, dữ liệu lưu ở MongoDB Atlas của UrBox. Nhà cung cấp AI đổi được, ví "
               "dụ sang Gemini. Mã nguồn có ba trăm bảy mươi sáu bài kiểm thử tự động."),
    # 10 -------------------------------------------------------------------------------------------------------
    dict(eyebrow="Nhớ xuyên cuộc họp", title="Biên bản tự nhắc khi buổi này nói khác điều đã chốt ở buổi trước", body=f"""
      <div class="split">
        {shot(IMG + '/14-bien-ban-thay-doi.jpg', 1000, 625)}
        <div class="rows">
          <div class="quote">"Dạ thưa anh chị, theo như mình đã bàn ở buổi 25/09/2026 thì là 15/12/2026 chứ không phải 20/12/2026, anh chị có thể xem xét lại nha ạ."</div>
          {row('history', 'So với 12 buổi gần nhất', 'Quyết định, con số, hạn chót, người phụ trách của nhóm.')}
          {row('file', 'Chỉ ghi vào biên bản', 'Không ngắt lời trong lúc họp; ghi rõ buổi nào.')}
          {row('users', 'Chủ nhóm hỏi lại bất cứ lúc nào', 'Trên toàn bộ cuộc họp của nhóm, có nguồn từng ý.')}
        </div>
      </div>""",
         hidden=["Biên bản nhắc điều nói khác buổi trước", "So với 12 buổi gần nhất của nhóm", "Chỉ ghi vào biên bản",
                 "Chủ nhóm hỏi trên toàn bộ cuộc họp"],
         message="Không ai phải tự nhớ đã chốt gì ở buổi nào.",
         qa=[("Có nhắc sai không, ai quyết định cuối cùng?",
              "AI so các điều đã chốt và ghi lại kèm buổi họp để người đọc xem xét; trợ lý không tự sửa quyết định.")],
         notes="Em nhớ xuyên các cuộc họp trong cùng một nhóm. Khi lập biên bản, em so với những điều nhóm đã chốt ở tối đa "
               "mười hai buổi trước: quyết định, con số, hạn chót, người phụ trách. Nếu buổi này nói khác, ví dụ ngày phát hành "
               "voucher là hai mươi tháng mười hai trong khi buổi trước đã chốt mười lăm, em ghi rõ trong biên bản kèm buổi họp "
               "nào. Em không ngắt lời trong lúc họp, chỉ ghi lại để mọi người xem xét."),
    # 11 -------------------------------------------------------------------------------------------------------
    dict(eyebrow="Trợ lý giọng nói", title="Gọi trợ lý bằng giọng nói tiếng Việt, tra cứu ngay trong buổi họp",
         body=f"""
      <div class="split rev">
        <div class="bubbles">
          <div class="bub">"Jarvis ơi, tóm tắt các quyết định từ đầu buổi"</div>
          <div class="bub">"Liệt kê việc cần làm, người phụ trách và hạn chót"</div>
          <div class="bub">"Tra cứu ticket Jira liên quan tới lỗi hoàn tiền"</div>
          <div class="bub">"Tra cứu trên mạng tỷ giá USD hôm nay"</div>
          <p class="cap">Jira hiện là dữ liệu mẫu qua chuẩn MCP; nối Jira thật là bước tiếp theo.</p>
        </div>
        {shot(IMG + '/08-tro-ly-tom-tat.jpg', 1000, 625)}
      </div>""",
         hidden=["Gọi Jarvis bằng giọng nói tiếng Việt", "Tóm tắt, liệt kê việc cần làm", "Tra cứu Jira qua MCP",
                 "Tra cứu trên mạng"],
         message="Không cần rời cuộc họp để tra cứu hay ghi chú.",
         qa=[("Nối được Jira thật không, cần gì?",
              "Trợ lý gọi công cụ qua MCP; hiện dùng dữ liệu mẫu. Nối Jira thật cần quyền đọc Jira và một MCP server cho "
              "Jira, cần team Engineering hỗ trợ.")],
         notes="Trong lúc họp, anh chị gọi em bằng tiếng Việt như gọi một đồng nghiệp: Jarvis ơi, tóm tắt các quyết định từ "
               "đầu buổi, hay liệt kê việc cần làm và người phụ trách. Em trả lời bằng giọng nói ngay, vừa làm vừa báo em đang "
               "làm gì. Em tra được thông tin trên mạng và dữ liệu nội bộ qua chuẩn MCP. Hiện phần Jira đang dùng dữ liệu mẫu; "
               "nối Jira thật là bước tiếp theo, cần sự hỗ trợ của team Engineering."),
    # 12 -------------------------------------------------------------------------------------------------------
    dict(eyebrow="Phần cơ bản", title="Phần cơ bản giống các công cụ khác, kèm vài tiện ích bổ sung", body=f"""
      <div class="cols2 basics">
        <div class="card">
          <div class="ch muted">Cơ bản, công cụ khác cũng có</div>
          <ul class="ticks two">
            <li>Chép lời</li><li>Tóm tắt</li><li>Biên bản Word</li><li>Việc cần làm</li>
            <li>Bot Google Meet theo lịch</li><li>Lưu Google Drive</li><li>Nhóm cuộc họp</li><li>Hỏi đáp trên các cuộc họp</li>
          </ul>
        </div>
        <div class="card white">
          <div class="ch">Tiện ích bổ sung</div>
          <ul class="ticks">
            <li>Soạn slide, dashboard, sơ đồ tư duy bằng một câu nói</li>
            <li>Mở tệp PDF, PowerPoint và tự trình bày, như bộ slide này</li>
          </ul>
          <div class="thumbs">{shot(D + '/p-slide.png', 150, 211)}{shot(D + '/p-dash.png', 150, 211)}{shot(D + '/p-mindmap.png', 150, 211)}</div>
        </div>
      </div>""",
         hidden=["Cơ bản: chép lời, tóm tắt, biên bản, bot Google Meet, lưu Drive",
                 "Tiện ích: slide, dashboard, sơ đồ tư duy, tự trình bày tệp"],
         message="Phần cơ bản có đủ; trọng tâm vẫn là hỗ trợ trong lúc họp.",
         qa=[("Các tiện ích như slide, dashboard có cần không?",
              "Không phải trọng tâm, có sẵn khi cần. Bộ slide đang xem là ví dụ: trợ lý mở tệp PowerPoint và tự trình bày "
              "theo ghi chú.")],
         notes="Phần còn lại là những tính năng cơ bản, giống các công cụ khác: chép lời, tóm tắt, biên bản Word, việc cần "
               "làm, bot vào Google Meet theo lịch, lưu Google Drive. Ngoài ra em có vài tiện ích bổ sung: soạn slide, dựng "
               "dashboard, vẽ sơ đồ tư duy bằng một câu nói, và mở tệp PowerPoint để tự trình bày, như chính bộ slide anh chị "
               "đang xem. Đây không phải trọng tâm, nhưng có sẵn khi cần."),
    # 13 -------------------------------------------------------------------------------------------------------
    dict(eyebrow="Dữ liệu và bảo mật", title="Biết rõ dữ liệu đi đâu, ai xem được", body=f"""
      <div class="cols2 data">
        <div class="card">
          <div class="ch">Dữ liệu gửi ra ngoài</div>
          <div class="ext"><b>Soniox</b><span>Âm thanh để chép lời; câu trợ lý nói để đọc thành tiếng</span></div>
          <div class="ext"><b>Anthropic</b><span>Nội dung cần AI xử lý (lời nói, tài liệu liên quan)</span></div>
          <div class="ext"><b>Recall</b><span>Âm thanh Google Meet khi có bot, xóa sau 1 giờ</span></div>
          <div class="ext"><b>Google</b><span>Biên bản trên Drive, lịch của chính người dùng</span></div>
        </div>
        <div class="tiles">
          <div class="t">{ic('lock', 26, box=52)}<span>Chỉ tài khoản @urbox.vn đăng nhập</span></div>
          <div class="t">{ic('users', 26, box=52)}<span>Chỉ người tạo và thành viên nhóm xem cuộc họp</span></div>
          <div class="t">{ic('mic', 26, box=52)}<span>Ghi âm cần mọi người đồng ý, tự xóa sau 30 ngày</span></div>
          <div class="t">{ic('shield', 26, box=52)}<span>Mẫu giọng chỉ lưu khi đồng ý, chỉ là vector</span></div>
          <div class="t">{ic('eyeoff', 26, box=52)}<span>Bảng BD không lên màn hình chung</span></div>
          <div class="t">{ic('file', 26, box=52)}<span>Tài liệu nhóm chỉ giữ phần chữ</span></div>
        </div>
      </div>""",
         hidden=["Gửi ra ngoài: Soniox, Anthropic, Recall, Google", "Chỉ tài khoản @urbox.vn",
                 "Ghi âm và mẫu giọng cần đồng ý", "Bảng BD không lên màn hình chung"],
         message="Luồng dữ liệu minh bạch, có đồng ý rõ ràng, phạm vi xem theo nhóm.",
         qa=[("Có cần review bảo mật trước khi dùng rộng không?",
              "Có. Đề xuất team Engineering cùng IT, Pháp chế review luồng dữ liệu ra ngoài và điều khoản của Soniox, "
              "Anthropic, Recall trước khi mở rộng.")],
         notes="Về dữ liệu, em liệt kê rõ những gì gửi ra ngoài: âm thanh tới Soniox để chép lời, nội dung cần AI xử lý tới "
               "Anthropic, âm thanh Google Meet qua Recall và được xóa sau một giờ, biên bản lưu ở Google Drive của chính người "
               "dùng. Chỉ tài khoản Google của công ty đăng nhập được, mỗi cuộc họp chỉ người tạo và thành viên nhóm xem được. "
               "Ghi âm và mẫu giọng chỉ lưu khi có đồng ý. Em đề xuất các anh chị review luồng dữ liệu này trước khi mở rộng."),
    # 14 -------------------------------------------------------------------------------------------------------
    dict(eyebrow="Bước tiếp theo", title="Đề xuất: thử ở team Engineering và team BD bằng cuộc họp thật", body=f"""
      <div class="cols3">
        <div class="card"><div class="ch">{ic('check', 26, '#1F8F55', 48)} Đã chạy</div>
          <ul><li>Bản web, đăng nhập Google</li><li>376 test tự động</li><li>BD và phỏng vấn thử với AI thật</li>
          <li>Nhận giọng đo trên 2 cuộc họp thật</li></ul></div>
        <div class="card"><div class="ch">{ic('clock', 26, '#B07600', 48)} Chưa kiểm chứng</div>
          <ul><li>Drive, Lịch, bot Meet với tài khoản thật</li><li>Nhận giọng trên nhiều cuộc họp, nhiều loại mic</li>
          <li>Nhiều phòng họp cùng lúc</li></ul></div>
        <div class="card white"><div class="ch">{ic('users', 26, box=48)} Cần từ Engineering</div>
          <ul><li>Review bảo mật và hạ tầng</li><li>Hỗ trợ nối Jira thật qua MCP</li>
          <li>Cho thử ở sprint review và phỏng vấn tuyển dụng</li></ul></div>
      </div>
      <div class="callout start">Bắt đầu nhỏ: 2 buổi sprint review, 2 buổi phỏng vấn, 3 buổi họp khách của BD, rồi đo lại độ chính xác và chi phí</div>""",
         hidden=["Đã chạy: bản web, 376 test, BD và phỏng vấn", "Chưa kiểm chứng: Drive, Lịch, bot Meet thật",
                 "Cần từ Engineering: review bảo mật, nối Jira, thử ở sprint review và phỏng vấn"],
         message="Sẵn sàng thử thật, bắt đầu nhỏ; cần Engineering đồng hành ở bảo mật, Jira và thử nghiệm.",
         qa=[("Muốn team Engineering làm gì cụ thể?",
              "Ba việc ở cột phải: review bảo mật và hạ tầng, hỗ trợ nối Jira thật, cho thử ở sprint review và các buổi "
              "phỏng vấn tuyển dụng."),
             ("Ai bảo trì sau khi thử nghiệm?",
              "Hiện do phuc.bh phát triển; mã nguồn có test và tài liệu, team Engineering có thể tham gia review từng bản.")],
         notes="Hiện bản web đã chạy, có ba trăm bảy mươi sáu bài kiểm thử, chế độ BD và phỏng vấn đã thử với AI thật. Những "
               "phần chưa kiểm chứng gồm lưu Google Drive, lịch và bot Google Meet với tài khoản thật, độ chính xác nhận giọng "
               "trên nhiều cuộc họp hơn, và khi nhiều phòng họp chạy cùng lúc. Em đề xuất thử ở team Engineering và team BD bằng "
               "các cuộc họp thật. Từ team Engineering, em cần ba việc: review bảo mật và hạ tầng, hỗ trợ nối Jira thật, và cho "
               "em thử ở các buổi sprint review và phỏng vấn tuyển dụng."),
    # 15 -------------------------------------------------------------------------------------------------------
    dict(kind="hero", title="Demo trực tiếp", body=f"""
      <div class="cl">
        <span class="pill-gold">Phần tiếp theo</span>
        <h1>Demo trực tiếp</h1>
        <ol class="agenda"><li>Họp với khách: khách hỏi giá, gợi ý có nguồn, tra thêm trên mạng</li>
          <li>Phỏng vấn: gợi ý câu hỏi đào sâu</li><li>Nhận ra người nói, người mới tự giới thiệu</li>
          <li>Gọi Jarvis tóm tắt, kết thúc xem biên bản</li></ol>
        <p class="gold">Sau đó là phần hỏi đáp</p>
      </div>
      <div class="cr"><div class="halo"></div><img class="mascot" src="{MASCOT}"></div>""",
         hidden=["Demo: họp với khách", "Demo: phỏng vấn", "Demo: nhận ra người nói", "Demo: tóm tắt và biên bản",
                 "Hỏi đáp"],
         message="Xem tận mắt theo đúng thứ tự ưu tiên: BD, phỏng vấn, nhận giọng, rồi phần cơ bản.",
         qa=[],
         notes="Phần trình bày của em đến đây là hết. Bây giờ em mời anh Phúc demo trực tiếp, theo đúng thứ tự ưu tiên: họp "
               "với khách trước, rồi phỏng vấn, nhận ra người nói, cuối cùng là tóm tắt và biên bản. Em cảm ơn anh chị đã lắng "
               "nghe."),
]

CSS = """
@import url('https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700;800&display=swap');
@page { size: 1920px 1080px; margin: 0; }
* { box-sizing: border-box; }
body { margin: 0; background: #fff; }
.slide { width: 1920px; height: 1080px; position: relative; overflow: hidden; font-family: 'Be Vietnam Pro', sans-serif;
  color: #27272A; background: #FFFFFF; page-break-after: always; }
.slide.light::before { content: ''; position: absolute; width: 1000px; height: 1000px; right: -420px; top: -560px; border-radius: 50%;
  background: radial-gradient(circle, rgba(130,53,228,.13), rgba(130,53,228,0) 68%); }
.slide.light::after { content: ''; position: absolute; width: 700px; height: 700px; left: -330px; bottom: -430px; border-radius: 50%;
  background: radial-gradient(circle, rgba(255,183,0,.10), rgba(255,183,0,0) 68%); }
.eyebrow { position: absolute; left: 110px; top: 70px; padding: 9px 22px; border-radius: 999px; background: #F1E9FC; color: #6821C2;
  font-size: 22px; font-weight: 700; z-index: 2; }
.title { position: absolute; left: 110px; top: 128px; width: 1700px; margin: 0; font-size: 50px; line-height: 1.22; font-weight: 800;
  color: #191919; letter-spacing: -0.5px; z-index: 2; text-wrap: balance; }
.content { position: absolute; left: 110px; right: 110px; top: 312px; bottom: 110px; z-index: 2; }
.foot { position: absolute; left: 110px; bottom: 40px; display: flex; align-items: center; gap: 14px; z-index: 2; }
.logo { height: 30px; display: block; }
.foot span { font-size: 20px; color: #6B6B76; font-weight: 600; }
.ic { border-radius: 50%; background: #F1E9FC; display: inline-grid; place-items: center; flex: none; }
.card { background: #F5F5FA; border-radius: 20px; padding: 36px 40px; }
.card.white { background: #fff; box-shadow: 0 14px 40px rgba(104,33,194,.12); border: 1px solid #EEE6FA; }
.ch { font-size: 30px; font-weight: 800; color: #191919; margin-bottom: 24px; display: flex; align-items: center; gap: 14px; }
.ch.muted { color: #5A5A66; }
.rows { display: flex; flex-direction: column; gap: 28px; flex: 1; }
.rows.tight { gap: 22px; }
.row { display: flex; gap: 22px; align-items: flex-start; }
.rh { font-size: 28px; font-weight: 700; color: #191919; line-height: 1.3; }
.rt { font-size: 23px; color: #4A4A55; line-height: 1.45; margin-top: 4px; }
.shot { border-radius: 20px; overflow: hidden; flex: none; background: #140E24;
  box-shadow: 0 24px 60px rgba(104,33,194,.22), 0 0 0 1px rgba(104,33,194,.10); }
.shot img { display: block; width: 100%; height: 100%; object-fit: cover; }
.split { display: flex; align-items: center; gap: 64px; height: 100%; }
.split .rows, .split .flowv, .split .bubbles { flex: 1; }
.cols2 { display: grid; grid-template-columns: 1fr 1fr; gap: 40px; height: 100%; }
.cols2.problem { grid-template-columns: 1fr 1.35fr; }
.chips { display: flex; flex-wrap: wrap; gap: 12px; margin-bottom: 26px; }
.chips span { background: #fff; border: 1px solid #E3DCEF; border-radius: 999px; padding: 8px 18px; font-size: 21px; font-weight: 600;
  color: #4A4A55; }
.chips.bottom { position: absolute; left: 0; bottom: 0; }
.chips.bottom span { background: #F1E9FC; border: 0; color: #6821C2; }
.ticks { list-style: none; padding: 0; margin: 0; }
.ticks li { font-size: 26px; line-height: 1.4; padding-left: 40px; margin-bottom: 16px; position: relative; color: #27272A; }
.ticks li::before { content: '✓'; position: absolute; left: 0; color: #6821C2; font-weight: 800; }
.ticks.two { display: grid; grid-template-columns: 1fr 1fr; column-gap: 20px; }
.foot-note { font-size: 22px; color: #6B6B76; margin: 18px 0 0; }
.diff { display: flex; flex-direction: column; gap: 14px; }
.drow { display: flex; gap: 26px; align-items: center; background: #F5F5FA; border-radius: 20px; padding: 18px 30px; }
.drow.top { background: #fff; box-shadow: 0 14px 40px rgba(104,33,194,.14); border: 1px solid #EEE6FA; }
.dn { width: 58px; height: 58px; border-radius: 50%; background: #6821C2; color: #fff; font-size: 28px; font-weight: 800;
  display: grid; place-items: center; flex: none; }
.drow.top .dn { background: #FFB700; color: #27272A; }
.dh { font-size: 28px; font-weight: 800; color: #191919; display: flex; align-items: center; gap: 16px; }
.dt { font-size: 22px; color: #4A4A55; margin-top: 4px; }
.badge-gold { background: #FFB700; color: #27272A; font-size: 19px; font-weight: 700; border-radius: 999px; padding: 5px 14px; }
.badge-gold.sm { font-size: 17px; padding: 3px 12px; }
.bottom-note { position: absolute; left: 0; bottom: 0; margin: 0; font-size: 22px; color: #6B6B76; }
.flowv { display: flex; flex-direction: column; gap: 22px; }
.step .sl { font-size: 22px; font-weight: 700; color: #6821C2; margin-bottom: 10px; display: flex; gap: 12px; align-items: center; }
.bubble { font-size: 27px; line-height: 1.4; background: #F5F5FA; border-radius: 20px 20px 20px 6px; padding: 20px 26px; font-weight: 600; }
.answer { font-size: 27px; line-height: 1.45; background: #fff; border-radius: 20px; padding: 22px 26px; font-weight: 600; color: #191919;
  box-shadow: 0 14px 40px rgba(104,33,194,.14); border: 1px solid #EEE6FA; }
.answer .note { font-size: 21px; font-weight: 500; color: #4A4A55; margin-top: 12px; }
.plain { font-size: 25px; line-height: 1.45; color: #27272A; }
.cite { display: inline-block; font-size: 16px; font-weight: 800; color: #6821C2; background: #F1E9FC; border-radius: 6px; padding: 2px 7px;
  vertical-align: 3px; }
.flowh { display: grid; grid-template-columns: repeat(4, 1fr); gap: 26px; }
.fcard { background: #F5F5FA; border-radius: 20px; padding: 32px 30px; position: relative; min-height: 300px; }
.fcard .fn { position: absolute; right: 26px; top: 18px; font-size: 56px; font-weight: 800; color: rgba(104,33,194,.16); }
.fh { font-size: 29px; font-weight: 800; color: #191919; margin: 22px 0 10px; }
.ft { font-size: 22px; color: #4A4A55; line-height: 1.45; }
.cols2.lower { height: auto; margin-top: 30px; }
.cols2.lower .card { display: flex; gap: 22px; align-items: center; padding: 28px 34px; }
.leftcol { display: flex; flex-direction: column; gap: 16px; }
.cap { font-size: 21px; color: #6B6B76; margin: 0; max-width: 960px; }
.nw { white-space: nowrap; }
.callout { font-size: 24px; font-weight: 700; color: #27272A; background: #FFF3CC; border-radius: 16px; padding: 16px 22px; }
.vflow { display: flex; align-items: stretch; gap: 14px; }
.vnode { flex: 1; background: #F5F5FA; border-radius: 20px; padding: 22px 24px; display: grid; grid-template-columns: 64px 1fr;
  column-gap: 18px; align-items: center; }
.vnode .ic { grid-row: span 2; }
.vnode.hl { background: #fff; box-shadow: 0 14px 40px rgba(104,33,194,.14); border: 1px solid #EEE6FA; }
.vh { font-size: 25px; font-weight: 800; color: #191919; }
.vt { font-size: 19px; color: #4A4A55; }
.varrow, .parrow { align-self: center; font-size: 46px; color: #8235E4; }
.vbottom { display: grid; grid-template-columns: 500px 1fr 1fr; gap: 34px; margin-top: 30px; }
.pair { display: flex; align-items: center; gap: 12px; }
.pair figure { margin: 0; }
.pair figcaption { font-size: 20px; font-weight: 700; text-align: center; margin-top: 10px; color: #4A4A55; }
.steps { margin: 0; padding-left: 30px; }
.steps li { font-size: 23px; line-height: 1.45; margin-bottom: 14px; color: #27272A; }
.stat .big { font-size: 66px; font-weight: 800; color: #6821C2; line-height: 1.1; }
.stat .rt { font-size: 23px; margin: 10px 0 20px; }
.mini { display: flex; gap: 12px; align-items: center; font-size: 21px; color: #27272A; font-weight: 600; }
.stats { display: grid; grid-template-columns: repeat(3, 1fr); grid-template-rows: 1fr 1fr; gap: 28px; height: calc(100% - 60px); }
.tile { background: #F5F5FA; border-radius: 20px; padding: 30px 34px; display: flex; flex-direction: column; justify-content: center; }
.tile .v { font-size: 52px; font-weight: 800; color: #6821C2; margin: 16px 0 6px; }
.tile .l { font-size: 23px; color: #4A4A55; line-height: 1.4; }
.arch { display: grid; grid-template-columns: 380px 120px 1fr 90px 470px; align-items: center; height: calc(100% - 70px); }
.acol { display: flex; flex-direction: column; gap: 22px; }
.acol.right { gap: 12px; }
.abox { display: flex; gap: 16px; align-items: center; background: #F5F5FA; border-radius: 18px; padding: 16px 20px; }
.abox b { display: block; font-size: 23px; color: #191919; }
.abox span { display: block; font-size: 18px; color: #4A4A55; }
.aarrow { text-align: center; font-size: 46px; color: #8235E4; line-height: 1; }
.aarrow span { display: block; font-size: 18px; font-weight: 700; color: #6821C2; margin-bottom: 4px; }
.acenter { background: #fff; border-radius: 24px; padding: 34px 38px; box-shadow: 0 18px 50px rgba(104,33,194,.16); border: 1px solid #EEE6FA; }
.atitle { font-size: 32px; font-weight: 800; color: #191919; margin-top: 18px; }
.asub { font-size: 21px; color: #6821C2; font-weight: 600; margin: 4px 0 16px; }
.acenter ul { margin: 0; padding-left: 26px; }
.acenter li { font-size: 22px; line-height: 1.45; margin-bottom: 8px; }
.quote { font-size: 25px; line-height: 1.45; font-style: italic; color: #3B1A6E; background: #F1E9FC; border-radius: 20px; padding: 24px 28px; }
.bubbles { display: flex; flex-direction: column; gap: 20px; }
.bub { font-size: 27px; line-height: 1.35; padding: 20px 28px; border-radius: 22px 22px 22px 6px; background: #F1E9FC; color: #3B1A6E; font-weight: 600; }
.basics .card { height: 100%; }
.thumbs { display: flex; gap: 22px; margin-top: 24px; }
.data { grid-template-columns: 1fr 1fr; }
.ext { display: flex; flex-direction: column; padding: 16px 0; border-bottom: 1px solid #E6E1F0; }
.ext:last-child { border-bottom: 0; }
.ext b { font-size: 26px; color: #6821C2; }
.ext span { font-size: 22px; color: #27272A; margin-top: 2px; }
.tiles { display: grid; grid-template-columns: 1fr 1fr; gap: 20px; }
.tiles .t { background: #F5F5FA; border-radius: 18px; padding: 22px 24px; display: flex; gap: 16px; align-items: center; font-size: 22px;
  font-weight: 600; line-height: 1.35; }
.cols3 { display: grid; grid-template-columns: repeat(3, 1fr); gap: 32px; }
.callout.start { margin-top: 34px; font-size: 26px; padding: 22px 28px; }
.cols3 ul { margin: 0; padding-left: 28px; }
.cols3 li { font-size: 25px; line-height: 1.4; margin-bottom: 18px; color: #27272A; }
/* bìa và slide demo: dải chuyển màu tím như urbox.vn */
.slide.hero { background: linear-gradient(118deg, #5A16AD 0%, #6821C2 38%, #8235E4 72%, #9BABFF 130%); color: #fff; }
.slide.hero::before { content: ''; position: absolute; width: 1200px; height: 1200px; right: -380px; top: -300px; border-radius: 50%;
  background: rgba(255,255,255,.06); }
.slide.hero::after { content: ''; position: absolute; width: 520px; height: 520px; left: -200px; bottom: -260px; border-radius: 50%;
  background: rgba(255,255,255,.05); }
.hero-logo { position: absolute; left: 110px; top: 76px; height: 54px; z-index: 2; }
.hero-body { position: absolute; inset: 0; display: flex; align-items: center; padding: 60px 110px 0; z-index: 2; }
.cl { flex: 1; }
.hero h1 { font-size: 100px; line-height: 1.04; margin: 30px 0 26px; font-weight: 800; letter-spacing: -2px; }
.sub { text-wrap: balance; font-size: 34px; line-height: 1.4; color: rgba(255,255,255,.92); margin: 0 0 34px; max-width: 1060px; }
.gold { font-size: 28px; color: #FFD166; font-weight: 700; margin: 0 0 16px; }
.who { font-size: 23px; color: rgba(255,255,255,.75); margin: 0; }
.pill-gold { display: inline-block; padding: 11px 26px; border-radius: 999px; background: #FFB700; color: #27272A; font-weight: 700; font-size: 24px; }
.cr { width: 600px; height: 600px; position: relative; flex: none; }
.halo { position: absolute; inset: 30px; border-radius: 50%; background: radial-gradient(circle, rgba(255,255,255,.30), rgba(255,255,255,0) 68%); }
.mascot { position: absolute; inset: 60px; width: 480px; height: 480px; object-fit: contain; }
.agenda { font-size: 32px; line-height: 1.5; margin: 0 0 30px; padding-left: 44px; color: rgba(255,255,255,.95); }
.agenda li { padding-left: 8px; margin-bottom: 6px; }
"""


def page_html():
    out = []
    logo = (HERE / "img" / "urbox-logo-white.png").as_uri()
    plogo = (HERE / "img" / "urbox-logo-purple.png").as_uri()
    for i, s in enumerate(SLIDES, 1):
        if s.get("kind") == "hero":
            out.append(f'<section class="slide hero" id="s{i}"><img class="hero-logo" src="{logo}">'
                       f'<div class="hero-body">{s["body"]}</div></section>')
        else:
            out.append(f'<section class="slide light" id="s{i}"><span class="eyebrow">{s["eyebrow"]}</span>'
                       f'<h2 class="title">{s["title"]}</h2><div class="content">{s["body"]}</div>'
                       f'<div class="foot"><img class="logo" src="{plogo}"><span>Meeting Copilot</span></div></section>')
    css = CSS
    return f'<!doctype html><html lang="vi"><head><meta charset="utf-8"><style>{css}</style></head><body>{"".join(out)}</body></html>'


def build():
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu, Pt
    work = Path(tempfile.mkdtemp(prefix="deck-"))
    html = work / "deck.html"
    html.write_text(page_html(), encoding="utf-8")
    pngs = []
    with sync_playwright() as p:
        b = p.chromium.launch()
        page = b.new_page(viewport={"width": 1920, "height": 1080})
        page.goto(html.as_uri())
        page.wait_for_load_state("networkidle")
        page.evaluate("document.fonts.ready")
        page.wait_for_timeout(800)
        for i in range(1, len(SLIDES) + 1):
            png = work / f"slide-{i:02d}.png"
            page.locator(f"#s{i}").screenshot(path=str(png))
            pngs.append(png)
        page.pdf(path=str(HERE / f"{NAME}.pdf"), width="1920px", height="1080px", print_background=True)
        b.close()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for s, png in zip(SLIDES, pngs):
        jpg = png.with_suffix(".jpg")
        Image.open(png).convert("RGB").save(jpg, quality=88, optimize=True)
        sl = prs.slides.add_slide(prs.slide_layouts[5])                      # Title Only: tiêu đề cho dàn ý, trình đọc màn hình
        sl.shapes.title.text = s["title"]
        tb = sl.shapes.add_textbox(Emu(600000), Emu(1500000), Emu(10900000), Emu(4500000))  # chữ ẩn dưới ảnh: tìm kiếm, nhắc bài
        tf = tb.text_frame
        tf.word_wrap = True
        for k, line in enumerate(s["hidden"]):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.text = line
            para.runs[0].font.size = Pt(14)
        pic = sl.shapes.add_picture(str(jpg), 0, 0, prs.slide_width, prs.slide_height)
        pic._element.nvPicPr.cNvPr.set("descr", s["title"] + ": " + "; ".join(s["hidden"]))
        sl.notes_slide.notes_text_frame.text = s["notes"]
    prs.core_properties.title = "UrBox Meeting Copilot - Demo"
    prs.core_properties.author = "phuc.bh@urbox.vn"
    prs.save(str(REPO / "slides" / "demo" / f"{NAME}.pptx"))
    for k, png in enumerate(pngs, 1):                                       # ảnh xem nhanh khi sửa (không đưa vào repo)
        png.replace(work / f"view-{k:02d}.png")
    write_notes_md()
    words = sum(len(s["notes"].split()) for s in SLIDES)
    print(json.dumps({"slides": len(SLIDES), "notes_words": words, "est_minutes": round(words / 140, 1),
                      "preview": str(work)}, ensure_ascii=False))


def _wrap(text: str, first: str = "", cont: str = "", width: int = 110) -> str:
    lines, cur = [], first
    for w in text.split():
        if cur.strip() and len(cur) + 1 + len(w) > width:
            lines.append(cur)
            cur = cont + w
        else:
            cur = f"{cur} {w}" if cur.strip() else cur + w
    return "\n".join(lines + [cur])


def write_notes_md():
    """Mục 5 của kịch bản demo: thông điệp, câu hỏi anh Tân có thể đặt ra, lời Jarvis (đúng ghi chú trong tệp)."""
    parts = []
    for i, s in enumerate(SLIDES, 1):
        p = [f"### Slide {i}. {s['title']}", "", _wrap(f"**Thông điệp:** {s['message']}")]
        if s["qa"]:
            p += ["", "**Anh Tân có thể hỏi:**", ""]
            for q, a in s["qa"]:
                p.append(_wrap(f"- *{q}* {a}", cont="  "))
        p += ["", "**Lời Jarvis:**", "", _wrap(s["notes"])]
        parts.append("\n".join(p))
    head = SCRIPT_MD.read_text(encoding="utf-8").split(MARK)[0]
    SCRIPT_MD.write_text(head + MARK + "\n\n" + "\n\n".join(parts) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build()
