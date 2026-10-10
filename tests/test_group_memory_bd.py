"""Giai đoạn 3 và 4 (bản 3.19): nhớ xuyên cuộc họp trong nhóm (mục "Thay đổi so với các buổi trước" trong biên bản), tài
liệu của nhóm, tìm kiếm BM25, bảng BD dạng khung chat (tự nhận câu hỏi của khách, trả lời gộp tài liệu + các buổi trước
có đánh dấu nguồn, kênh riêng), trò chuyện của chủ nhóm (đọc song song các cuộc họp).
AI là bản giả: không gọi Claude / Gemini thật."""
import asyncio
import io
import json
import os
import tempfile
import time
import unittest
from unittest import mock

from fastapi.testclient import TestClient

from tests.helpers import reset_db
from meeting import app as appmod
from meeting import artifacts, auth, bd, db, group_chat, group_kb, group_memory, groups, live, retrieval

WEB = dict(ENABLED=True, SECRET=b"k" * 32, DOMAINS=["urbox.vn"], ADMINS=set(),
           CLIENT_ID="123-test.apps.googleusercontent.com", TTL_S=3600)
DAY = 86400
T0 = 1_780_000_000.0
MINUTES_OLD = """# BIÊN BẢN CUỘC HỌP: Review voucher Tết
## 1. Tóm Tắt Điều Hành
Nhóm thống nhất kế hoạch phát hành voucher Tết.
## 2. Các Quyết Định Quan Trọng Đã Thống Nhất
- Ngày phát hành voucher Tết: 15/12/2026.
- Chị Lan phụ trách đối soát với đối tác.
"""


def hdr(email):
    return {"Authorization": f"Bearer {auth.issue({'email': email, 'name': email})['token']}"}


def meeting(title, gid, started, minutes=None, lines=(), owner="an@urbox.vn", ended=True):
    mid = db.create_meeting(title, owner=owner)
    if gid:
        groups.set_meeting_group(mid, gid)
    db.update_meeting(mid, {"started_at": started, "status": "ended" if ended else "live",
                            "ended_at": started + 3600 if ended else None}, True)
    for i, (spk, text) in enumerate(lines):
        db.add_segment(mid, 10.0 * i, 10.0 * i + 8, spk, text, seq=i + 1)
    if minutes:
        db.save_artifact(meeting_id=mid, kind="minutes", title=f"Biên bản: {title}", content=minutes, prompt_trigger="t")
    return mid


class FakeLLM:
    """AI giả: trả lời theo loại việc (system prompt), ghi lại lời gọi, đo số lời gọi chạy cùng lúc."""

    def __init__(self, replies, delay=0.0):
        self.replies, self.delay, self.calls = replies, delay, []
        self.active = self.max_active = 0

    async def __call__(self, system, prompt, max_tokens=4000):
        self.calls.append({"system": str(system), "prompt": str(prompt), "owner": artifacts.CURRENT_OWNER.get(),
                           "meeting": artifacts.CURRENT_MEETING.get(), "purpose": artifacts.CURRENT_PURPOSE.get()})
        self.active += 1
        self.max_active = max(self.max_active, self.active)
        try:
            if self.delay:
                await asyncio.sleep(self.delay)
            for key, reply in self.replies.items():
                if key in str(system):
                    return reply(str(prompt)) if callable(reply) else reply
            raise AssertionError(f"lời gọi không mong đợi: {str(system)[:80]}")
        finally:
            self.active -= 1

    def of(self, key):
        return [c for c in self.calls if key in c["system"]]


class RetrievalTests(unittest.TestCase):
    def test_vietnamese_with_or_without_accents_and_compound_words(self):
        self.assertIn("phat_hanh", retrieval.terms("Phát hành voucher"))
        idx = retrieval.Index([
            {"kind": "kb", "title": "Bảng giá", "text": "Phí tích hợp API: 5.000.000đ một lần, phí duy trì 500.000đ mỗi tháng."},
            {"kind": "kb", "title": "Chính sách", "text": "Voucher hết hạn được gia hạn một lần trong 30 ngày."},
            {"kind": "talk", "meeting_id": 7, "title": "Họp Vinmart", "text": "Khách hỏi phí tích hợp API bao nhiêu"}])
        top = idx.search("phi tich hop bao nhieu", k=3)                      # lời nói nhận dạng mất dấu
        self.assertEqual({p["title"] for p in top[:2]}, {"Bảng giá", "Họp Vinmart"})
        self.assertEqual([p["title"] for p in idx.search("gia hạn voucher", kinds={"kb"})][:1], ["Chính sách"])
        self.assertEqual(idx.search("phí tích hợp", exclude_meeting=7, kinds={"talk"}), [])
        self.assertEqual(idx.search("và là của"), [])                           # chỉ có từ phổ biến

    def test_minutes_sections_and_transcript_windows(self):
        secs = retrieval.minutes_sections(MINUTES_OLD)
        self.assertEqual([h for h, _ in secs], ["1. Tóm Tắt Điều Hành", "2. Các Quyết Định Quan Trọng Đã Thống Nhất"])
        segs = [{"t_start": i * 30, "speaker_label": "An", "text": "x" * 300} for i in range(5)]
        wins = retrieval.transcript_windows(segs, size=700)
        self.assertEqual(len(wins), 2)
        self.assertTrue(wins[1][1].startswith("[01:30] An:"))


class GroupDocsApiTests(unittest.TestCase):
    def setUp(self):
        reset_db()
        self.p = mock.patch.multiple(auth, **WEB)
        self.p.start()
        self.client = TestClient(appmod.app)
        self.client.__enter__()
        self.g = groups.create_group("BD Vinmart", "an@urbox.vn")
        groups.update_group(self.g["id"], members=["binh@urbox.vn"], bd_mode=True)

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.p.stop()

    def upload(self, name, data, who="an@urbox.vn"):
        return self.client.post(f"/api/groups/{self.g['id']}/docs", files={"file": (name, io.BytesIO(data))}, headers=hdr(who))

    def test_owner_uploads_member_reads_text_only_is_kept(self):
        md = ("# Bảng giá dịch vụ\n\n" + "\n\n".join(f"Gói {i}: phí tích hợp {i}.000.000đ, hỗ trợ 24/7." for i in range(1, 40))).encode()
        r = self.upload("bang-gia.md", md)
        self.assertEqual(r.status_code, 200, r.text)
        self.assertGreater(r.json()["n_chunks"], 1)
        import docx
        d = docx.Document()
        d.add_heading("FAQ đối tác", 1)
        d.add_paragraph("Thời gian đối soát: trước ngày 5 hằng tháng.")
        buf = io.BytesIO()
        d.save(buf)
        self.assertEqual(self.upload("faq.docx", buf.getvalue()).status_code, 200)
        self.assertEqual(self.upload("virus.exe", b"MZ").status_code, 400)
        self.assertEqual(self.upload("trong.md", b"   ").status_code, 400)
        self.assertEqual(self.upload("cua-binh.md", md, who="binh@urbox.vn").status_code, 403)       # chỉ chủ nhóm
        listing = self.client.get(f"/api/groups/{self.g['id']}/docs", headers=hdr("binh@urbox.vn")).json()
        self.assertEqual([x["name"] for x in listing["docs"]], ["faq.docx", "bang-gia.md"])
        self.assertFalse(listing["can_edit"])
        self.assertEqual(self.client.get(f"/api/groups/{self.g['id']}/docs", headers=hdr("la@urbox.vn")).status_code, 404)
        stored = db._get_db()["group_docs"].find_one({"name": "faq.docx"})
        self.assertIn("đối soát", " ".join(stored["chunks"]))
        self.assertNotIn("path", stored)                                         # không giữ tệp gốc
        texts = [p["text"] for p in group_kb.passages(self.g["id"])]
        self.assertTrue(any("Gói 39" in t for t in texts))
        did = listing["docs"][0]["id"]
        self.assertEqual(self.client.delete(f"/api/groups/{self.g['id']}/docs/{did}", headers=hdr("binh@urbox.vn")).status_code, 403)
        self.assertEqual(self.client.delete(f"/api/groups/{self.g['id']}/docs/{did}", headers=hdr("an@urbox.vn")).status_code, 200)
        info = self.client.get(f"/api/groups/{self.g['id']}", headers=hdr("an@urbox.vn")).json()
        self.assertEqual((info["bd_mode"], info["doc_count"]), (True, 1))
        groups.delete_group(self.g["id"])
        self.assertEqual(db._get_db()["group_docs"].count_documents({}), 0)    # xóa nhóm xóa luôn tài liệu

    def test_only_owner_switches_bd_mode(self):
        r = self.client.patch(f"/api/groups/{self.g['id']}", json={"bd_mode": False}, headers=hdr("binh@urbox.vn"))
        self.assertEqual(r.status_code, 403)
        r = self.client.patch(f"/api/groups/{self.g['id']}", json={"bd_mode": False}, headers=hdr("an@urbox.vn"))
        self.assertFalse(r.json()["bd_mode"])


class GroupMemoryTests(unittest.TestCase):
    def setUp(self):
        reset_db()
        self.g = groups.create_group("Sprint voucher", "an@urbox.vn")
        other = groups.create_group("Nhóm khác", "an@urbox.vn")
        self.old = meeting("Review voucher Tết", self.g["id"], T0 - 7 * DAY, minutes=MINUTES_OLD)
        self.older = meeting("Kickoff", self.g["id"], T0 - 14 * DAY, minutes="# Kickoff\n## 2. Quyết định\n- Ngân sách 200.000.000đ")
        db.update_meeting(self.old, {"memory": {"facts": [
            {"topic": "Ngày phát hành voucher Tết", "value": "15/12/2026", "kind": "hạn chót", "who": "An"},
            {"topic": "Người phụ trách đối soát", "value": "Chị Lan", "kind": "người phụ trách", "who": ""}],
            "changes": [], "at": T0 - 7 * DAY}}, True)
        meeting("Họp nhóm khác", other["id"], T0 - 3 * DAY, minutes="# Khác\n- Ngày phát hành 01/01/2027")
        self.mid = meeting("Review sprint 41", self.g["id"], T0, ended=False,
                           lines=[("An", "Mình dời phát hành voucher Tết sang 20/12 nhé"), ("Lan", "Dạ em đối soát như cũ")])

    def run_minutes(self, mid, fake):
        with mock.patch.object(artifacts, "_call_llm", fake):
            m = db.get_meeting(mid)
            return asyncio.run(artifacts.generate_meeting_minutes(mid, db.get_segments(mid), m["title"], meeting=m))

    def test_minutes_note_what_changed_from_earlier_meetings_of_the_group(self):
        new_minutes = "# BIÊN BẢN CUỘC HỌP: Review sprint 41\n## 2. Các Quyết Định\n- Phát hành voucher Tết ngày 20/12/2026."

        def memory(prompt):
            p = next(line for line in prompt.splitlines() if "Ngày phát hành voucher Tết" in line)
            code = p.split("]")[0].strip("[")
            return json.dumps({"facts": [{"topic": "Ngày phát hành voucher Tết", "value": "20/12/2026", "kind": "hạn chót"}],
                               "changes": [{"topic": "Ngày phát hành voucher Tết", "before": "15/12/2026", "ref": code,
                                            "now": "20/12/2026"},
                                           {"topic": "Bịa ra", "before": "x", "ref": "P99", "now": "y"}]}, ensure_ascii=False)
        fake = FakeLLM({"Thư ký Cuộc họp Thông minh": new_minutes,
                        "Đọc BIÊN BẢN và trả về": json.dumps({"facts": [{"topic": "Ngân sách", "value": "200.000.000đ",
                                                                           "kind": "con số"}]}, ensure_ascii=False),
                        "CÁC ĐIỀU ĐÃ CHỐT Ở CÁC BUỔI TRƯỚC, mỗi dòng": memory})
        art = self.run_minutes(self.mid, fake)
        date = retrieval.fmt_date(T0 - 7 * DAY)
        self.assertIn("## Thay Đổi So Với Các Buổi Trước", art["content"])
        self.assertIn(f"Dạ thưa anh chị, theo như mình đã bàn ở buổi {date} (Review voucher Tết) thì là 15/12/2026 chứ "
                      "không phải 20/12/2026, anh chị có thể xem xét lại nha ạ.", art["content"])
        self.assertNotIn("Bịa ra", art["content"])                              # mã không có trong danh sách: bỏ
        self.assertEqual(db.get_artifact(art["id"])["content"], art["content"])  # lưu cả mục thay đổi (Drive, Word có)
        mem = db.get_meeting(self.mid)["memory"]
        self.assertEqual(mem["facts"][0]["value"], "20/12/2026")
        self.assertEqual(mem["changes"][0]["ref"], self.old)
        self.assertEqual(len(fake.of("Đọc BIÊN BẢN và trả về")), 1)              # rút bù buổi Kickoff chưa có
        self.assertEqual(db.get_meeting(self.older)["memory"]["facts"][0]["topic"], "Ngân sách")
        prompt = fake.of("CÁC ĐIỀU ĐÃ CHỐT Ở CÁC BUỔI TRƯỚC, mỗi dòng")[0]["prompt"]
        self.assertIn("Ngân sách: 200.000.000đ", prompt)
        self.assertNotIn("01/01/2027", prompt)                                  # nhóm khác không lẫn vào
        self.assertEqual(fake.of("CÁC ĐIỀU ĐÃ CHỐT")[0]["purpose"], "nhớ xuyên cuộc họp")

    def test_no_change_no_section_and_ungrouped_meeting_untouched(self):
        fake = FakeLLM({"Thư ký Cuộc họp Thông minh": "# Biên bản\n- Giữ ngày 15/12/2026",
                        "Đọc BIÊN BẢN và trả về": '{"facts": []}',
                        "CÁC ĐIỀU ĐÃ CHỐT Ở CÁC BUỔI TRƯỚC, mỗi dòng": '{"facts": [], "changes": []}'})
        art = self.run_minutes(self.mid, fake)
        self.assertNotIn("Thay Đổi", art["content"])
        solo = meeting("Họp riêng", None, T0, ended=False, lines=[("An", "Dời sang 20/12")])
        fake2 = FakeLLM({"Thư ký Cuộc họp Thông minh": "# Biên bản riêng"})
        self.assertEqual(self.run_minutes(solo, fake2)["content"], "# Biên bản riêng")
        self.assertEqual(len(fake2.calls), 1)


class BDSessionTests(unittest.TestCase):
    def setUp(self):
        reset_db()
        self.g = groups.create_group("BD Vinmart", "an@urbox.vn")
        groups.update_group(self.g["id"], bd_mode=True)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "bang-gia.md")
            with open(path, "w", encoding="utf-8") as f:
                f.write("# Bảng giá\n\nPhí tích hợp API: 5.000.000đ một lần.\n\nHỗ trợ kỹ thuật 24/7 qua hotline.")
            group_kb.add(self.g["id"], "bang-gia.md", path, "an@urbox.vn", 100)
        self.prev = meeting("Họp Vinmart lần 1", self.g["id"], T0 - 7 * DAY,
                            lines=[("Bình", "Lần trước bên em báo phí tích hợp API là 4.000.000đ cho Vinmart")])
        self.mid = meeting("Họp Vinmart lần 2", self.g["id"], T0, ended=False)
        self.patches = [mock.patch.object(bd, "DEBOUNCE_S", 0.01), mock.patch.object(bd, "MIN_GAP_S", 0.3)]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in self.patches:
            p.stop()

    def fake(self, question="Phí tích hợp API bao nhiêu?", delay=0.05):
        def reply(prompt):
            self.assertIn("Phí tích hợp API: 5.000.000đ", prompt)          # nguồn tài liệu [K..]
            self.assertIn("4.000.000đ", prompt)                              # nguồn cuộc họp trước [M..]
            return json.dumps({"question": question,
                               "answer": "Dạ phí tích hợp API bên em là 5.000.000đ một lần ạ [K1].",
                               "notes": ["Buổi trước anh Bình đã báo 4.000.000đ [M1], nên giải thích phần chênh lệch."],
                               "confidence": "cao"}, ensure_ascii=False)
        return FakeLLM({"Bạn là trợ lý ngồi cạnh đội kinh doanh (BD)": reply}, delay=delay)

    def test_customer_question_is_answered_in_the_chat_from_documents_and_earlier_meetings(self):
        fake = self.fake()

        async def go():
            s = await live.get_session(self.mid)
            room = await s.subscribe()                                          # kênh chung (màn hình trình chiếu)
            a = bd.assistant(s)
            q = await a.subscribe()                                             # kênh riêng của bảng BD
            with mock.patch.object(artifacts, "_call_llm", fake):
                await s.on_segment_finalized(1.0, 3.0, "1", "Dạ vâng", stream="mic")
                await s.drain()
                await asyncio.sleep(0.1)
                n_after_ok = len(fake.calls)
                await s.on_segment_finalized(4.0, 7.0, "2", "Vậy phí tích hợp API bên em là bao nhiêu?", stream="mic")
                await s.drain()
                for _ in range(100):
                    await asyncio.sleep(0.02)
                    if fake.calls and a.inflight == 0 and not a._tasks - {a._timer}:
                        break
                await a.drain()
            events = []
            while not q.empty():
                events.append(q.get_nowait())
            room_types = []
            while not room.empty():
                room_types.append(room.get_nowait()["type"])
            return n_after_ok, events, room_types, s
        n_after_ok, events, room_types, s = asyncio.run(go())
        self.assertEqual(n_after_ok, 0)                                         # "Dạ vâng": không gọi AI
        self.assertEqual(len(fake.calls), 1)                                    # một câu hỏi, một câu trả lời gộp
        turns = [e["turn"] for e in events if e["type"] == "bd_turn"]
        self.assertEqual([t["status"] for t in turns], ["pending", "done"])     # câu hỏi hiện ngay, trả lời tới sau
        self.assertEqual(turns[0]["question"], "Vậy phí tích hợp API bên em là bao nhiêu?")
        done = turns[-1]
        self.assertEqual((done["source"], done["question"]), ("auto", "Phí tích hợp API bao nhiêu?"))
        self.assertEqual(done["seq"], s.segments[-1]["seq"])                    # đánh dấu câu trong lời nói
        self.assertIn("[K1]", done["answer"])
        self.assertEqual([x["code"] for x in done["sources"]], ["K1", "M1"])
        self.assertEqual(done["sources"][0]["title"], "bang-gia.md")
        self.assertEqual(done["sources"][1]["meeting_id"], self.prev)
        self.assertIn("[M1]", done["notes"][0])
        self.assertNotIn("bd_turn", room_types)                                 # màn hình trình chiếu không nhận gì
        self.assertEqual([t["status"] for t in bd.load_turns(self.mid)], ["done"])
        self.assertEqual({c["purpose"] for c in fake.calls}, {"BD gợi ý"})

    def test_not_a_real_question_is_removed_from_the_chat(self):
        fake = self.fake(question="")

        async def go():
            s = await live.get_session(self.mid)
            a = bd.assistant(s)
            q = await a.subscribe()
            with mock.patch.object(artifacts, "_call_llm", fake):
                await s.on_segment_finalized(1.0, 3.0, "1", "Bên em có phí tích hợp API riêng đúng không, em nhớ là vậy", stream="mic")
                await s.drain()
                for _ in range(100):
                    await asyncio.sleep(0.02)
                    if fake.calls and a.inflight == 0:
                        break
                await a.drain()
            return [q.get_nowait() for _ in range(q.qsize())]
        events = asyncio.run(go())
        statuses = [e["turn"]["status"] for e in events if e["type"] == "bd_turn"]
        self.assertEqual(statuses, ["pending", "none"])
        self.assertEqual(bd.load_turns(self.mid), [])

    def test_team_lines_ignored_follow_up_questions_and_marked_lines(self):
        fake = self.fake(delay=0.3)

        async def go():
            s = await live.get_session(self.mid)
            a = bd.assistant(s)
            with mock.patch.object(artifacts, "_call_llm", fake), mock.patch.object(bd, "DEBOUNCE_S", 30):
                await s.on_segment_finalized(1.0, 3.0, "1", "Bên em có hỗ trợ không ạ?", stream="mic")
                await s.drain()
                seg = s.segments[-1]
                await a.set_role(seg["speaker_key"], "team")
                a.last_seq = 0
                await a.check()                                                 # câu của đội mình: không tự trả lời
                await asyncio.sleep(0.05)
                team_calls = len(fake.calls)
                first = await a.ask("Phí tích hợp API bao nhiêu?", "binh@urbox.vn")
                await a.drain()
                both = [await a.ask("Còn hỗ trợ ngoài giờ thì sao?", "binh@urbox.vn"),
                        await a.mark(seg["seq"], "an@urbox.vn")]              # hai câu hỏi chạy song song
                await a.drain()
                a.close()
            return team_calls, first, both
        team_calls, first, both = asyncio.run(go())
        self.assertEqual(team_calls, 0)
        self.assertEqual((first["source"], first["by"]), ("typed", "binh@urbox.vn"))
        follow = fake.calls[1]["prompt"]
        self.assertIn("HỘI THOẠI TRƯỚC TRÊN KHUNG CHAT", follow)               # hiểu câu hỏi nối tiếp
        self.assertIn("Phí tích hợp API bao nhiêu?", follow)
        self.assertEqual(fake.max_active, 2)
        self.assertEqual((both[1]["source"], both[1]["question"]), ("marked", "Bên em có hỗ trợ không ạ?"))
        self.assertEqual([t["source"] for t in bd.load_turns(self.mid)], ["typed", "typed", "marked"])
        self.assertEqual(list(db.get_meeting(self.mid)["bd"]["roles"].values()), ["team"])

    def test_web_lookup_runs_only_when_asked_and_cites_web_sources(self):
        def web_reply(prompt):
            self.assertIn("CÂU TRẢ LỜI TỪ TÀI LIỆU NỘI BỘ: Dạ phí tích hợp API bên em là 5.000.000đ", prompt)
            self.assertNotIn("[K1]", prompt.split("CÂU TRẢ LỜI TỪ TÀI LIỆU NỘI BỘ:")[1].split("\n")[0])
            self.assertIn("[W1] Bảng giá đối thủ - doithu.vn", prompt)
            self.assertIn("Phí tích hợp 3.000.000đ một lần.", prompt)
            return json.dumps({"answer": "Một đơn vị khác niêm yết phí tích hợp 3.000.000đ [W1].",
                               "notes": ["Giá của UrBox vẫn theo bảng giá nội bộ; trang thứ hai không liên quan [W2]."],
                               "confidence": "vừa"}, ensure_ascii=False)
        base = self.fake()
        fake = FakeLLM({'Đội BD vừa bấm "Tra thêm': web_reply, **base.replies})
        searched = []

        async def fake_web(query, on_progress=None, rewrite=True, max_uses=5, basic=False):
            self.assertEqual((max_uses, basic), (bd.W_SEARCHES, True))   # BD: tìm nhanh, có trích dẫn
            searched.append(query)
            await asyncio.sleep(0.05)
            if len(searched) == 2:                                              # lượt tra thứ hai: công cụ tìm kiếm lỗi
                return {"error": "bị chặn"}
            return {"query": query, "engine": "Bing", "sources": [
                {"n": 1, "title": "Bảng giá đối thủ", "url": "https://doithu.vn/bang-gia", "domain": "doithu.vn",
                 "published": "", "excerpt": "Phí tích hợp 3.000.000đ một lần."},
                {"n": 2, "title": "Trang lạ", "url": "javascript:alert(1)", "domain": "la.vn", "excerpt": "không liên quan"},
                {"n": 3, "title": "Không dùng", "url": "https://khac.vn", "domain": "khac.vn", "excerpt": "..."}]}

        async def go():
            s = await live.get_session(self.mid)
            a = bd.assistant(s)
            q = await a.subscribe()
            with mock.patch.object(artifacts, "_call_llm", fake), mock.patch.object(artifacts, "web_search_tool", fake_web):
                t = await a.ask("Phí tích hợp API bao nhiêu?", "binh@urbox.vn")
                with self.assertRaises(ValueError):                             # đang trả lời: chưa tra trên mạng
                    await a.search_web(t["id"], "binh@urbox.vn")
                await a.drain()
                before = list(searched)
                first = await a.search_web(t["id"], "binh@urbox.vn")
                await a.search_web(t["id"], "binh@urbox.vn")                    # bấm hai lần: chỉ tra một lần
                await a.drain()
                with self.assertRaises(KeyError):
                    await a.search_web("khong-co", None)
                t2 = await a.ask("Câu hỏi làm lỗi tra cứu?", None)
                await a.drain()
                await a.search_web(t2["id"], None)
                await a.drain()
            return before, first, [q.get_nowait() for _ in range(q.qsize())]
        before, first, events = asyncio.run(go())
        self.assertEqual(before, [])                                            # trả lời thường không tra trên mạng
        self.assertEqual(searched, ["Phí tích hợp API bao nhiêu?", "Phí tích hợp API bao nhiêu?"])
        self.assertEqual((first["web"]["status"], first["web"]["by"]), ("pending", "binh@urbox.vn"))
        turns = bd.load_turns(self.mid)
        web = turns[0]["web"]
        self.assertEqual(web["status"], "done")
        self.assertIn("[W1]", web["answer"])
        self.assertEqual([(x["code"], x["domain"], x["url"]) for x in web["sources"]],
                         [("W1", "doithu.vn", "https://doithu.vn/bang-gia"), ("W2", "la.vn", "")])   # bỏ link javascript:
        self.assertEqual(turns[0]["answer"], "Dạ phí tích hợp API bên em là 5.000.000đ một lần ạ [K1].")  # câu trả lời cũ giữ nguyên
        self.assertEqual((turns[1]["web"]["status"], turns[1]["web"]["error"]), ("error", "bị chặn"))
        webs = [e["turn"]["web"]["status"] for e in events if e["type"] == "bd_turn" and e["turn"].get("web")]
        self.assertEqual(webs, ["pending", "done", "pending", "error"])
        self.assertEqual({c["purpose"] for c in fake.of('Đội BD vừa bấm "Tra thêm')}, {"BD tra cứu web"})

    def test_web_answer_survives_truncated_json(self):
        base = self.fake()
        cut = '{"answer": "Tỷ giá \\"trung tâm\\" ngày 10/10/2026 là 25.629đ [W1].", "notes": ["Ghi chú dài bị cắt giữa chừ'
        fake = FakeLLM({'Đội BD vừa bấm "Tra thêm': cut, **base.replies})

        async def fake_web(query, on_progress=None, rewrite=True, **kw):
            return {"query": query, "engine": "Claude web_search", "summary": "Tỷ giá trung tâm 25.629đ.",
                    "sources": [{"n": 1, "title": "Tỷ giá", "url": "https://sbv.gov.vn", "domain": "sbv.gov.vn", "excerpt": ""}]}

        async def go():
            s = await live.get_session(self.mid)
            a = bd.assistant(s)
            with mock.patch.object(artifacts, "_call_llm", fake), mock.patch.object(artifacts, "web_search_tool", fake_web):
                t = await a.ask("Tỷ giá USD hôm nay?", None)
                await a.drain()
                await a.search_web(t["id"], None)
                await a.drain()
        asyncio.run(go())
        web = bd.load_turns(self.mid)[0]["web"]
        self.assertEqual((web["status"], web["answer"]), ("done", 'Tỷ giá "trung tâm" ngày 10/10/2026 là 25.629đ [W1].'))
        self.assertEqual([x["code"] for x in web["sources"]], ["W1"])

    def test_group_without_bd_mode_has_no_assistant(self):
        g2 = groups.create_group("Sprint", "an@urbox.vn")
        mid = meeting("Họp sprint", g2["id"], T0, ended=False)

        async def go():
            return bd.assistant(await live.get_session(mid))
        self.assertIsNone(asyncio.run(go()))


class BDApiTests(unittest.TestCase):
    def setUp(self):
        reset_db()
        self.p = mock.patch.multiple(auth, **WEB)
        self.p.start()
        self.client = TestClient(appmod.app)
        self.client.__enter__()
        self.g = groups.create_group("BD Vinmart", "an@urbox.vn")
        groups.update_group(self.g["id"], members=["binh@urbox.vn"], bd_mode=True)
        self.mid = meeting("Họp Vinmart", self.g["id"], T0, ended=False, lines=[("Khách", "Phí tích hợp bao nhiêu?")])
        plain = groups.create_group("Sprint", "an@urbox.vn")
        self.plain = meeting("Họp sprint", plain["id"], T0, ended=False)

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.p.stop()

    def test_bd_panel_api_and_private_socket(self):
        from starlette.websockets import WebSocketDisconnect
        r = self.client.get(f"/api/meetings/{self.mid}/bd", headers=hdr("binh@urbox.vn")).json()
        self.assertTrue(r["enabled"])
        self.assertEqual((r["group"]["name"], r["turns"]), ("BD Vinmart", []))
        self.assertFalse(self.client.get(f"/api/meetings/{self.plain}/bd", headers=hdr("an@urbox.vn")).json()["enabled"])
        self.assertEqual(self.client.get(f"/api/meetings/{self.mid}/bd", headers=hdr("la@urbox.vn")).status_code, 404)
        tok = lambda e: hdr(e)["Authorization"][7:]
        with self.client.websocket_connect(f"/ws/meeting/{self.mid}/bd?token={tok('binh@urbox.vn')}") as ws:
            self.assertEqual(json.loads(ws.receive_text())["type"], "bd_init")
        with self.assertRaises(WebSocketDisconnect):
            with self.client.websocket_connect(f"/ws/meeting/{self.plain}/bd?token={tok('an@urbox.vn')}") as ws:
                self.assertEqual(json.loads(ws.receive_text())["type"], "error")
                ws.receive_text()
        with self.assertRaises(WebSocketDisconnect):                           # người ngoài nhóm
            with self.client.websocket_connect(f"/ws/meeting/{self.mid}/bd?token={tok('la@urbox.vn')}") as ws:
                ws.receive_text()
        r = self.client.put(f"/api/meetings/{self.mid}/bd/roles/3", json={"role": "client"}, headers=hdr("binh@urbox.vn"))
        self.assertEqual(r.json()["roles"], {"3": "client"})
        self.assertEqual(self.client.put(f"/api/meetings/{self.mid}/bd/roles/3", json={"role": "boss"},
                                         headers=hdr("binh@urbox.vn")).status_code, 400)

        async def no_answer(self_, turn):
            return None
        with mock.patch.object(bd.BDAssistant, "_answer", no_answer):
            r = self.client.post(f"/api/meetings/{self.mid}/bd/ask", json={"question": "Phí bao nhiêu?"},
                                 headers=hdr("binh@urbox.vn"))
            self.assertEqual(r.status_code, 200, r.text)
            self.assertEqual((r.json()["turn"]["status"], r.json()["turn"]["by"]), ("pending", "binh@urbox.vn"))
            self.assertEqual(self.client.post(f"/api/meetings/{self.mid}/bd/mark", json={"seq": 99},
                                              headers=hdr("binh@urbox.vn")).status_code, 404)
            r = self.client.post(f"/api/meetings/{self.mid}/bd/mark", json={"seq": 1}, headers=hdr("binh@urbox.vn"))
            self.assertEqual((r.json()["turn"]["source"], r.json()["turn"]["question"]), ("marked", "Phí tích hợp bao nhiêu?"))
        self.assertEqual(len(self.client.get(f"/api/meetings/{self.mid}/bd", headers=hdr("an@urbox.vn")).json()["turns"]), 2)
        tid = r.json()["turn"]["id"]
        self.assertEqual(self.client.post(f"/api/meetings/{self.mid}/bd/web", json={"turn_id": tid},
                                          headers=hdr("binh@urbox.vn")).status_code, 409)       # câu chưa trả lời xong
        self.assertEqual(self.client.post(f"/api/meetings/{self.mid}/bd/web", json={"turn_id": "x"},
                                          headers=hdr("binh@urbox.vn")).status_code, 404)
        self.assertEqual(self.client.post(f"/api/meetings/{self.mid}/bd/web", json={"turn_id": tid},
                                          headers=hdr("la@urbox.vn")).status_code, 404)          # người ngoài nhóm
        db.delete_meeting(self.mid)
        self.assertEqual(db._get_db()["bd_turns"].count_documents({}), 0)       # xóa cuộc họp xóa luôn khung chat BD


class GroupChatTests(unittest.TestCase):
    def setUp(self):
        reset_db()
        self.g = groups.create_group("BD Vinmart", "an@urbox.vn")
        groups.update_group(self.g["id"], members=["binh@urbox.vn"], bd_mode=True)
        self.mids = []
        for i in range(10):
            lines = [("Khách", f"Buổi {i}: trao đổi tiến độ chung")]
            if i in (3, 6):
                lines.append(("Khách", "Bên chị phàn nàn đối soát chậm, tháng nào cũng trễ hạn"))
            self.mids.append(meeting(f"Họp Vinmart {i}", self.g["id"], T0 + i * DAY, lines=lines,
                                     minutes=f"# Buổi {i}\n## 1. Tóm tắt\nTrao đổi tiến độ."))

    def test_picks_relevant_and_recent_meetings(self):
        picked = [m["id"] for m in group_chat.pick_meetings(self.g["id"], "Khách phàn nàn gì về đối soát?")]
        self.assertLessEqual(len(picked), group_chat.MAX_MEETINGS)
        for mid in (self.mids[3], self.mids[6], self.mids[9], self.mids[8]):  # 2 buổi liên quan + 2 buổi gần nhất
            self.assertIn(mid, picked)

    def test_reads_meetings_in_parallel_and_answers_with_sources(self):
        def mapper(prompt):
            return "- Khách phàn nàn đối soát chậm [00:10]" if "đối soát chậm" in prompt else "KHÔNG CÓ"

        def reducer(prompt):
            self.assertIn("2 cuộc họp có thông tin liên quan", prompt)
            return "Khách phàn nàn đối soát chậm ở 2 buổi [M1] [M2]."
        fake = FakeLLM({"Bạn đọc MỘT cuộc họp": mapper, "Bạn là trợ lý phân tích cho chủ nhóm": reducer}, delay=0.05)

        async def go():
            with mock.patch.object(artifacts, "_call_llm", fake):
                msg = group_chat.prepare(self.g["id"], "Báo cáo các câu phàn nàn của khách", "an@urbox.vn")
                await group_chat.start(self.g["id"], msg, "an@urbox.vn")
            return msg
        msg = asyncio.run(go())
        done = next(m for m in group_chat.messages(self.g["id"]) if m["id"] == msg["id"])
        self.assertEqual(done["status"], "done")
        self.assertIn("[M1]", done["text"])
        self.assertEqual(sorted(s["meeting_id"] for s in done["sources"]), sorted([self.mids[3], self.mids[6]]))
        self.assertEqual(len(fake.of("Bạn đọc MỘT cuộc họp")), group_chat.MAX_MEETINGS)
        self.assertGreater(fake.max_active, 1)                                  # đọc song song
        self.assertLessEqual(fake.max_active, group_chat.PARALLEL)
        self.assertEqual({c["owner"] for c in fake.calls}, {"an@urbox.vn"})   # chạy bằng gói AI riêng của chủ nhóm
        self.assertEqual([m["role"] for m in group_chat.messages(self.g["id"])], ["user", "assistant"])

    def test_chat_api_owner_only_bd_groups_only(self):
        with mock.patch.multiple(auth, **WEB):
            with TestClient(appmod.app) as client:
                url = f"/api/groups/{self.g['id']}/chat"
                self.assertEqual(client.get(url, headers=hdr("binh@urbox.vn")).status_code, 403)
                with mock.patch.object(group_chat, "start") as start:
                    r = client.post(url, json={"question": "Khách đang gặp vấn đề gì?"}, headers=hdr("an@urbox.vn"))
                self.assertEqual(r.status_code, 200, r.text)
                start.assert_called_once()
                self.assertEqual(r.json()["message"]["status"], "pending")
                self.assertEqual(client.post(url, json={"question": "Câu khác"}, headers=hdr("an@urbox.vn")).status_code, 400)
                self.assertEqual(client.delete(url, headers=hdr("an@urbox.vn")).json()["messages"], [])
                groups.update_group(self.g["id"], bd_mode=False)
                self.assertEqual(client.get(url, headers=hdr("an@urbox.vn")).status_code, 400)


if __name__ == "__main__":
    unittest.main()
