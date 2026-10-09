"""Xuất biên bản (Google Docs / Word) và ghi âm (MP3) để lưu vào thư mục cuộc họp trên Drive và trên máy."""
import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import numpy as np

from meeting import recap_export, recording

MINUTES = """# BIÊN BẢN CUỘC HỌP: Review Sprint 40
**Thời gian:** 09/10/2026 14:00 | **Chủ trì:** Bùi Hồng Phúc | **Thành viên tham dự:** An, Bình

## 1. Tóm Tắt Điều Hành (Executive Summary)
Nhóm thống nhất phát hành **voucher Tết** trước ngày 15/12.
Còn vướng phần đối soát với đối tác.

## 2. Các Quyết Định Quan Trọng Đã Thống Nhất (Key Decisions)
- **Phát hành ngày 15/12**: đã chốt, anh Phúc đề xuất.
- Giữ nguyên giá 1.000.000đ

## 3. Bảng Phân Công Nhiệm Vụ (Action Items Matrix)
| STT | Nhiệm vụ (Action Item) | Người phụ trách (Assignee) | Hạn chót (Deadline) |
|---|---|---|---|
| 1 | Đối soát với đối tác | An | 20/10/2026 |
| 2 | Thiết kế banner <Tết> | Bình | 25/10/2026 |

## 4. Vấn Đề Tồn Đọng / Rủi Ro Cần Theo Dõi (Open Questions & Risks)
1. Đối tác chưa gửi file đối soát
2. Rủi ro trễ hạn
"""


class MarkdownExportTests(unittest.TestCase):
    def test_blocks_cover_the_minutes_structure(self):
        kinds = [k for k, _ in recap_export.md_blocks(MINUTES)]
        self.assertEqual(kinds, ["h1", "p", "h2", "p", "h2", "ul", "h2", "table", "h2", "ol"])
        table = dict(recap_export.md_blocks(MINUTES))["table"]
        self.assertEqual(table[0], ["STT", "Nhiệm vụ (Action Item)", "Người phụ trách (Assignee)", "Hạn chót (Deadline)"])
        self.assertEqual(len(table), 3)                      # dòng |---| bị bỏ

    def test_html_for_google_docs(self):
        html = recap_export.md_to_html(MINUTES, "Biên bản - Review Sprint 40")
        for tag in ("<h1>", "<h2>", "<table", "<th>", "<li>", "<b>voucher Tết</b>", "<ol>", "<ul>"):
            self.assertIn(tag, html)
        self.assertNotIn("**", html)
        self.assertIn("&lt;Tết&gt;", html)                   # ký tự HTML trong biên bản được thoát
        self.assertIn('<meta charset="utf-8">', html)

    def test_word_document(self):
        import docx
        d = docx.Document(io.BytesIO(recap_export.md_to_docx_bytes(MINUTES)))
        heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
        self.assertEqual(heads[0], "BIÊN BẢN CUỘC HỌP: Review Sprint 40")
        self.assertEqual(len(d.tables), 1)
        self.assertEqual(d.tables[0].cell(2, 1).text, "Thiết kế banner <Tết>")
        bold = [r.text for p in d.paragraphs for r in p.runs if r.bold]
        self.assertIn("voucher Tết", bold)
        bullets = [p.text for p in d.paragraphs if p.style.name == "List Bullet"]
        self.assertEqual(bullets, ["Phát hành ngày 15/12: đã chốt, anh Phúc đề xuất.", "Giữ nguyên giá 1.000.000đ"])

    def test_folder_name_uses_vietnam_time_and_safe_characters(self):
        ts = 1791529200                                       # 09/10/2026 07:00 UTC = 14:00 giờ Việt Nam
        name = recap_export.folder_name({"started_at": ts, "title": 'Họp: Q4/2026 "chốt" <giá>?'})
        self.assertEqual(name, "09-10-2026 14h00 - Họp- Q4-2026 -chốt- -giá-")
        self.assertEqual(recap_export.folder_name({"started_at": ts, "title": ""}), "09-10-2026 14h00 - Cuộc họp")
        self.assertLessEqual(len(recap_export.folder_name({"started_at": ts, "title": "x" * 500})), 120)


class RecordingMp3Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.p = mock.patch.object(recording, "ROOT", Path(self.tmp.name))
        self.p.start()

    def tearDown(self):
        self.p.stop()
        self.tmp.cleanup()

    def test_runs_are_joined_into_one_mp3(self):
        import soundfile as sf
        tone = (np.sin(np.arange(16000) / 16000 * 2 * np.pi * 440) * 6000).astype("<i2").tobytes()
        for start in (0.0, 30.0):
            r = recording.Run(7, "mic", start)
            r.write(tone)
            r.close()
        path = recap_export.recording_mp3(7)
        try:
            info = sf.info(str(path))
            self.assertEqual(info.samplerate, 16000)
            self.assertAlmostEqual(info.duration, 2.0, delta=0.3)
        finally:
            recap_export.discard(path)
        self.assertFalse(path.exists())
        self.assertIsNone(recap_export.recording_mp3(8))     # cuộc họp không ghi âm


if __name__ == "__main__":
    unittest.main()


class RecapDownloadApiTests(unittest.TestCase):
    """Tải biên bản Word / ghi âm MP3 và danh sách biên bản đã xong (trình duyệt tự lưu về thư mục trên máy)."""

    def setUp(self):
        from tests.helpers import reset_db
        reset_db()
        self.tmp = tempfile.TemporaryDirectory()
        self.p = mock.patch.object(recording, "ROOT", Path(self.tmp.name))
        self.p.start()

    def tearDown(self):
        self.p.stop()
        self.tmp.cleanup()

    def test_docx_mp3_and_ready_list(self):
        import docx
        from fastapi.testclient import TestClient
        from meeting import app as appmod, db, groups
        g = groups.create_group("Sprint: Payment", None)
        mid = db.create_meeting("Review Sprint 40")
        groups.set_meeting_group(mid, g["id"])
        db.update_meeting(mid, {"status": "ended", "minutes_status": "done", "ended_at": 1791532800.0,
                                "started_at": 1791529200.0}, True)
        aid = db.save_artifact(meeting_id=mid, kind="minutes", title="Biên bản", content=MINUTES, prompt_trigger="t")
        other = db.create_meeting("Chưa có biên bản")
        r = recording.Run(mid, "mic", 0.0)
        r.write((np.sin(np.arange(16000) / 16000 * 2 * np.pi * 300) * 5000).astype("<i2").tobytes())
        r.close()
        with TestClient(appmod.app) as c:
            res = c.get(f"/api/meetings/{mid}/recap.docx")
            self.assertEqual(res.status_code, 200)
            self.assertIn("wordprocessingml", res.headers["content-type"])
            self.assertEqual(len(docx.Document(io.BytesIO(res.content)).tables), 1)
            mp3 = c.get(f"/api/meetings/{mid}/recording.mp3")
            self.assertEqual((mp3.status_code, mp3.headers["content-type"]), (200, "audio/mpeg"))
            self.assertGreater(len(mp3.content), 500)
            self.assertEqual(c.get(f"/api/meetings/{other}/recap.docx").status_code, 404)
            self.assertEqual(c.get(f"/api/meetings/{other}/recording.mp3").status_code, 404)
            items = c.get("/api/recaps/ready").json()["items"]
        self.assertEqual(len(items), 1)
        it = items[0]
        self.assertEqual((it["id"], it["minutes_id"], it["has_recording"]), (mid, aid, True))
        self.assertEqual((it["group_folder"], it["folder_name"]), ("Sprint- Payment", "09-10-2026 14h00 - Review Sprint 40"))
