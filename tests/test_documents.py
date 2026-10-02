"""Test mở tài liệu trên máy (PDF, PowerPoint, Word), hỏi "tự trình bày hay theo kịch bản", và nguồn AI qua gói đăng ký."""
import asyncio
import json
import os
import subprocess
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from fastapi.testclient import TestClient

from tests.helpers import reset_db
from meeting import app as appmod
from meeting import artifacts, cli_llm, db, decks, live, llm


def make_pdf(path: Path, pages):
    """PDF tối giản viết tay: mỗi trang vài dòng chữ Helvetica (pdfium đọc được chữ và dựng được ảnh)."""
    objs = {1: b"<< /Type /Catalog /Pages 2 0 R >>"}
    kids, n = [], 3
    font_id = 3 + 2 * len(pages)
    for lines in pages:
        ops = ["BT /F1 26 Tf 40 300 Td"] + [f"({ln}) Tj 0 -36 Td" for ln in lines] + ["ET"]
        stream = "\n".join(ops).encode("latin-1")
        objs[n] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 640 360] /Contents {n + 1} 0 R "
                   f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>").encode()
        objs[n + 1] = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        kids.append(f"{n} 0 R")
        n += 2
    objs[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    objs[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    out, offsets = bytearray(b"%PDF-1.4\n"), {}
    for i in sorted(objs):
        offsets[i] = len(out)
        out += f"{i} 0 obj\n".encode() + objs[i] + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for i in sorted(objs):
        out += f"{offsets[i]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def make_pptx(path: Path, png: Path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Mục tiêu quý 4"
    s1.placeholders[1].text_frame.text = "Doanh số 2.500.000.000đ"
    s1.placeholders[1].text_frame.add_paragraph().text = "120 thương hiệu mới"
    s1.shapes.add_picture(str(png), Inches(6), Inches(2), Inches(3), Inches(2))
    s1.notes_slide.notes_text_frame.text = "Nhấn mạnh doanh số tăng 25% so với quý 3."
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Rủi ro"
    s2.placeholders[1].text_frame.text = "Banner mobile trễ 3 ngày"
    prs.save(str(path))


def make_docx(path: Path):
    import docx
    d = docx.Document()
    d.add_heading("Kế hoạch Mega Sale", 0)
    d.add_heading("Mục tiêu", 1)
    d.add_paragraph("Doanh số 2.500.000.000đ", style="List Bullet")
    d.add_paragraph("Tăng 120 thương hiệu", style="List Bullet")
    d.add_heading("Rủi ro", 1)
    d.add_paragraph("Banner mobile trễ", style="List Bullet")
    d.save(str(path))


class TempLibrary:
    """Thư mục tài liệu tạm: slides/ và ảnh trang đều nằm trong thư mục test."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name) / "slides"
        self.base.mkdir()
        self.assets = Path(self.tmp.name) / "assets"
        self.patches = [mock.patch.object(decks, "SLIDES_DIR", self.base), mock.patch.object(decks, "ASSETS_DIR", self.assets)]
        for p in self.patches:
            p.start()
        decks.invalidate()

    def tearDown(self):
        for p in self.patches:
            p.stop()
        decks.invalidate()
        self.tmp.cleanup()


class ReadFileTests(TempLibrary, unittest.TestCase):
    def test_pdf_pages_become_image_slides_with_text(self):
        pdf = self.base / "ke-hoach-q4.pdf"
        make_pdf(pdf, [["Ke hoach Q4", "Doanh so tang 25%"], ["Rui ro", "Banner mobile tre"]])
        deck = artifacts.normalize_deck(decks.read_deck(str(pdf)), max_slides=decks.MAX_PAGES)
        self.assertEqual([s["title"] for s in deck["slides"]], ["Ke hoach Q4", "Rui ro"])
        self.assertEqual(deck["slides"][0]["layout"], "image")
        self.assertIn("Doanh so tang 25%", deck["slides"][0]["bullets"])
        key, name = deck["slides"][1]["image"].split("/")[-2:]
        self.assertEqual(name, "2.jpg")
        self.assertIsNotNone(decks.asset_path(key, name))                   # ảnh trang đã dựng trên máy
        self.assertEqual(deck["source"], {"type": "file", "path": str(pdf), "name": "ke-hoach-q4.pdf", "ext": ".pdf"})
        self.assertIn("Banner mobile tre", decks.read_text(str(pdf)))

    def test_pptx_text_notes_and_picture(self):
        from PIL import Image
        png = Path(self.tmp.name) / "chart.png"
        Image.new("RGB", (300, 200), (61, 220, 151)).save(png)
        pp = self.base / "mega-sale.pptx"
        make_pptx(pp, png)
        with mock.patch.object(decks, "_soffice", lambda: None):
            deck = artifacts.normalize_deck(decks.read_deck(str(pp)), max_slides=decks.MAX_PAGES)
        s1, s2 = deck["slides"]
        self.assertEqual((s1["title"], s1["layout"]), ("Mục tiêu quý 4", "media"))
        self.assertEqual(s1["bullets"], ["Doanh số 2.500.000.000đ", "120 thương hiệu mới"])
        self.assertEqual(s1["notes"], "Nhấn mạnh doanh số tăng 25% so với quý 3.")
        self.assertRegex(s1["image"], r"/api/deck-assets/[0-9a-f]{16}/1m\.png$")
        self.assertEqual((s2["layout"], s2.get("image")), ("bullets", None))
        self.assertTrue(decks.read_text(str(pp)).startswith("Slide 1: Nhấn mạnh doanh số"))

    def test_docx_headings_become_slides(self):
        dx = self.base / "ke-hoach.docx"
        make_docx(dx)
        deck = artifacts.normalize_deck(decks.read_deck(str(dx)))
        self.assertEqual(deck["title"], "Kế hoạch Mega Sale")
        self.assertEqual([s["title"] for s in deck["slides"]], ["Mục tiêu", "Rủi ro"])
        self.assertEqual(deck["slides"][0]["bullets"], ["Doanh số 2.500.000.000đ", "Tăng 120 thương hiệu"])

    def test_bad_asset_names_are_refused(self):
        self.assertIsNone(decks.asset_path("../../etc", "1.jpg"))
        self.assertIsNone(decks.asset_path("0123456789abcdef", "..\\x.jpg"))
        self.assertIsNone(artifacts.normalize_deck({"title": "x", "slides": [{"title": "a", "layout": "image",
                                                                              "image": "https://evil.example/a.png"}]})["slides"][0].get("image"))

    def test_only_files_inside_library_can_be_opened(self):
        inside = self.base / "a.md"
        inside.write_text("# A\n## B\n- c", encoding="utf-8")
        outside = Path(self.tmp.name) / "secret.txt"
        outside.write_text("mật khẩu", encoding="utf-8")
        self.assertEqual(decks.allowed(str(inside)), inside.resolve())
        self.assertIsNone(decks.allowed(str(outside)))
        self.assertIsNone(decks.allowed(str(self.base / "khong-co.pdf")))
        self.assertEqual(decks.find_files(f"mở {outside}"), [])            # đường dẫn ngoài thư mục tài liệu: không mở


class ScriptSplitTests(unittest.TestCase):
    def test_split_by_marks_and_paragraphs(self):
        text = "Kính chào anh chị.\nSlide 1: Mục tiêu quý 4 là doanh số.\nSlide 2: Rủi ro lớn nhất là banner.\nSlide 3: Cảm ơn."
        self.assertEqual(decks.split_script(text, 3),
                         ["Kính chào anh chị.\nMục tiêu quý 4 là doanh số.", "Rủi ro lớn nhất là banner.", "Cảm ơn."])
        self.assertEqual(decks.split_script("Đoạn một.\n\nĐoạn hai.", 2), ["Đoạn một.", "Đoạn hai."])
        self.assertIsNone(decks.split_script("Đoạn một.\n\nĐoạn hai.\n\nĐoạn ba.", 2))
        self.assertEqual(decks.split_script("Chỉ một trang.", 1), ["Chỉ một trang."])

    def test_present_answers(self):
        cases = {
            "em tự trình bày đi": {"mode": "auto"},
            "không cần kịch bản": {"mode": "auto"},
            "để anh trình bày": {"mode": "human"},
            "theo kịch bản trong ghi chú": {"mode": "script", "source": "notes"},
            "theo kịch bản": {"mode": "script", "source": None},
            "anh dán kịch bản vào khung chat": {"mode": "script", "source": "text"},
            "em tự soạn slide về doanh thu": None,
            "tóm tắt giúp anh": None,
        }
        for text, want in cases.items():
            self.assertEqual(llm.present_answer(text), want, text)
        self.assertEqual(llm.present_answer("kịch bản ở file kich-ban trong Downloads")["source"], "file")
        self.assertIsNone(llm.present_answer("trình bày đi"))                    # chưa hỏi: không tự hiểu là trả lời
        self.assertEqual(llm.present_answer("trình bày đi", asked=True), {"mode": "auto"})
        self.assertIsNone(llm.present_answer("mở file báo cáo trong Downloads", asked=True))
        # câu chung chung kèm lệnh: vẫn hỏi; đang chờ trả lời: nhận cả câu ngắn
        self.assertIsNone(llm.present_answer("em thuyết trình giúp anh", explicit=True))
        self.assertEqual(llm.present_answer("mở file mega rồi tự trình bày luôn", explicit=True), {"mode": "auto"})
        self.assertIsNone(llm.present_answer("mở cái file anh nói lúc nãy", explicit=True))
        self.assertIsNone(llm.present_answer("mở file báo cáo có sẵn", explicit=True))
        self.assertEqual(llm.present_answer("để anh", asked=True), {"mode": "human"})
        self.assertIsNone(llm.present_answer("để anh xem", asked=True))
        self.assertEqual(llm.stage_intent("mở file pdf kế hoạch q4")["action"], "open_file")
        self.assertEqual(llm.stage_intent("em mở cái tài liệu powerpoint mega sale")["action"], "open_file")


class PresentDialogTests(TempLibrary, unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):              # TempLibrary.setUp / tearDown chạy tự động trước / sau phần async
        reset_db()
        llm.set_assistant_config("Bông", [])
        self.mid = db.create_meeting("Họp")
        self.s = await live.get_session(self.mid)
        self.q = await self.s.subscribe()
        self.ev = []
        make_pdf(self.base / "ke-hoach-q4.pdf", [["Ke hoach Q4", "Doanh so tang 25%"], ["Rui ro", "Banner mobile tre"]])
        (self.base / "mega.md").write_text("# Mega Sale\n## Mục tiêu\n- Doanh số\nNhấn mạnh doanh số tăng.\n## Rủi ro\n- Banner",
                                           encoding="utf-8")

    async def asyncTearDown(self):
        for s in list(live.SESSIONS.values()):
            s.dispose()
        live.SESSIONS.clear()

    def events(self, kind=None):
        while not self.q.empty():
            self.ev.append(self.q.get_nowait())
        return [e for e in self.ev if kind is None or e["type"] == kind]

    async def wait_for(self, kind, timeout=4.0):
        for _ in range(int(timeout / 0.02)):
            if self.events(kind):
                return self.events(kind)
            await asyncio.sleep(0.02)
        self.fail(f"không thấy sự kiện {kind}: {[e['type'] for e in self.events()]}")

    async def open_pdf(self):
        await self.s._handle_ai_activation("mở file pdf kế hoạch q4", "Bông ơi, mở file pdf kế hoạch q4", "Bông")
        ask = self.events("ai_ask")
        self.assertEqual((ask[0]["kind"], ask[0]["has_notes"]), ("present_mode", False))
        return db.get_artifact(ask[0]["artifact_id"])

    async def test_open_pdf_asks_then_voice_answer_auto_writes_scripts(self):
        art = await self.open_pdf()
        deck = json.loads(art["content"])
        self.assertEqual((deck["slides"][0]["layout"], len(deck["slides"])), ("image", 2))
        says = [e["text"] for e in self.events("ai_say")]
        self.assertTrue(any("ke-hoach-q4.pdf" in t and "2 trang" in t for t in says), says)

        async def fake(system, prompt, max_tokens=4000):
            self.assertIn("LỜI THUYẾT TRÌNH", system)
            return json.dumps({"scripts": [("lời trình bày trang " * 10).strip()] * 2}, ensure_ascii=False)
        with mock.patch.object(artifacts, "_call_llm", fake), mock.patch.object(artifacts, "llm_available", lambda: True):
            # trả lời không cần gọi tên trợ lý, ngay sau câu hỏi
            await self.s.on_segment_finalized(t_start=1.0, t_end=2.5, raw_speaker="1", text="Em tự trình bày đi.",
                                              vector=None, voiced=1.2, epoch=0)
            await self.s.drain()
            await self.wait_for("stage_present")
        new = db.get_artifact(self.s.stage["artifact_id"])
        nd = json.loads(new["content"])
        self.assertEqual((new["parent_id"], nd["script_mode"]), (art["id"], "auto"))
        self.assertTrue(all(s["script"] for s in nd["slides"]))
        self.assertTrue(self.events("ai_ask_done"))
        self.assertIsNone(self.s._ask)

    async def test_answer_window_ignores_other_talk(self):
        await self.open_pdf()
        calls = []

        async def fake_activation(command, full_sentence, name="", source="voice"):
            calls.append(command)
        self.s._handle_ai_activation = fake_activation
        await self.s.on_segment_finalized(t_start=1.0, t_end=2.5, raw_speaker="2", text="Hôm nay trời mưa to quá.",
                                          vector=None, voiced=1.2, epoch=0)
        await self.s.drain()
        await asyncio.sleep(0)
        self.assertEqual(calls, [])
        await self.s.on_segment_finalized(t_start=3.0, t_end=4.5, raw_speaker="1", text="Theo ghi chú trong file nhé.",
                                          vector=None, voiced=1.2, epoch=0)
        await self.s.drain()
        await asyncio.sleep(0)
        self.assertEqual(calls, ["Theo ghi chú trong file nhé."])

    async def test_notes_script_is_read_verbatim(self):
        await self.s._handle_ai_activation("mở file mega", "Bông ơi, mở file mega", "Bông")
        self.assertTrue(self.events("ai_ask")[0]["has_notes"])
        res = await self.s.answer_present({"mode": "script", "source": "notes"})
        self.assertTrue(res["ok"])
        nd = json.loads(db.get_artifact(res["artifact_id"])["content"])
        self.assertEqual((nd["script_mode"], nd["slides"][0]["script"], nd["slides"][1]["script"]),
                         ("script", "Nhấn mạnh doanh số tăng.", ""))
        says = [e["text"] for e in self.events("ai_say")]
        self.assertTrue(any("theo ghi chú trong tệp" in t and "1 slide chưa có lời" in t for t in says), says)
        self.assertEqual([e["action"] for e in self.events("stage_present")], ["start"])

    async def test_no_notes_asks_where_then_file_answer(self):
        await self.open_pdf()
        res = await self.s.answer_present({"mode": "script", "source": "notes"})
        self.assertEqual(res.get("need"), "where")
        self.assertEqual(self.events("ai_ask")[-1]["kind"], "script_where")
        (self.base / "kich-ban-q4.txt").write_text("Trang 1: Chào anh chị, đây là kế hoạch quý 4.\nTrang 2: Rủi ro chính là banner.",
                                                   encoding="utf-8")
        decks.invalidate()
        ans = llm.present_answer("kịch bản ở file kich ban q4", asked=True)
        res = await self.s.answer_present(ans)
        self.assertTrue(res["ok"], res)
        nd = json.loads(db.get_artifact(res["artifact_id"])["content"])
        self.assertEqual([s["script"] for s in nd["slides"]], ["Chào anh chị, đây là kế hoạch quý 4.", "Rủi ro chính là banner."])
        self.assertEqual(nd["script_source"], "tệp kich-ban-q4.txt")

    async def test_pasted_script_without_marks_is_aligned_by_ai(self):
        await self.open_pdf()

        async def fake(system, prompt, max_tokens=4000):
            self.assertIn("ghép kịch bản", system)
            return json.dumps({"assign": [1, 1, 2]})
        with mock.patch.object(artifacts, "_call_llm", fake), mock.patch.object(artifacts, "llm_available", lambda: True):
            res = await self.s.answer_present({"mode": "script", "source": "text",
                                               "text": "Chào anh chị.\nĐây là kế hoạch quý 4.\nRủi ro là banner trễ."})
        nd = json.loads(db.get_artifact(res["artifact_id"])["content"])
        self.assertEqual([s["script"] for s in nd["slides"]], ["Chào anh chị.\nĐây là kế hoạch quý 4.", "Rủi ro là banner trễ."])

    async def test_generic_present_request_asks_for_file_deck(self):
        art = await self.open_pdf()
        self.s._ask = None                                                   # câu hỏi cũ đã hết hạn

        async def no_llm(*a, **k):
            raise AssertionError("không được gọi AI khi chưa chọn cách trình bày")
        with mock.patch.object(artifacts, "_call_llm", no_llm), mock.patch.object(artifacts, "llm_available", lambda: True):
            await self.s._handle_ai_activation("em thuyết trình giúp anh", "Bông ơi, em thuyết trình giúp anh", "Bông")
        asks = self.events("ai_ask")
        self.assertEqual((len(asks), asks[-1]["artifact_id"]), (2, art["id"]))
        self.assertEqual(self.events("stage_present"), [])

    async def test_generic_present_request_keeps_existing_scripts(self):
        deck = {"title": "Q4", "slides": [{"title": "A", "bullets": ["x"], "script": "lời " * 30},
                                          {"title": "B", "bullets": ["y"], "script": "lời " * 30}]}
        aid = db.save_artifact(self.mid, "slides", "Slide: Q4", json.dumps(deck, ensure_ascii=False))
        await self.s.stage_action("show", artifact_id=aid)

        async def no_llm(*a, **k):
            raise AssertionError("bộ slide đã có lời thuyết trình, không soạn lại")
        with mock.patch.object(artifacts, "_call_llm", no_llm), mock.patch.object(artifacts, "llm_available", lambda: True):
            await self.s._handle_ai_activation("em thuyết trình giúp anh", "Bông ơi, em thuyết trình giúp anh", "Bông")
        self.assertEqual([e["action"] for e in self.events("stage_present")], ["start"])
        self.assertEqual((self.events("ai_ask"), self.events("artifact_updated")), ([], []))

    async def test_human_presents_ai_follows(self):
        await self.open_pdf()
        res = await self.s.answer_present({"mode": "human"})
        self.assertEqual(res, {"ok": True, "mode": "human"})
        self.assertTrue(self.s.stage["follow"])
        self.assertTrue(self.events("ai_ask_done"))
        self.assertEqual(self.events("stage_present"), [])

    async def test_inline_answer_in_open_command(self):
        (self.base / "mega.md").write_text("# Mega\n## A\n- x\nGhi chú A.\n## B\n- y\nGhi chú B.", encoding="utf-8")
        await self.s._handle_ai_activation("mở file mega rồi trình bày theo ghi chú", "Bông ơi, mở file mega rồi trình bày theo ghi chú", "Bông")
        self.assertEqual(self.events("ai_ask"), [])                          # đã nói cách trình bày: không hỏi lại
        nd = json.loads(db.get_artifact(self.s.stage["artifact_id"])["content"])
        self.assertEqual([s["script"] for s in nd["slides"]], ["Ghi chú A.", "Ghi chú B."])


class CliProviderTests(unittest.TestCase):
    def setUp(self):
        self.env = mock.patch.dict(os.environ, {"CLI_LLM_DISABLED": "0", "ANTHROPIC_API_KEY": "sk-test", "OPENAI_API_KEY": "sk-o"})
        self.env.start()
        self.which = mock.patch.object(cli_llm, "resolve_cmd", lambda name: [f"/bin/{name}"])
        self.which.start()

    def tearDown(self):
        self.which.stop()
        self.env.stop()

    def test_claude_cli_uses_subscription_without_tools(self):
        seen = {}

        def fake_run(cmd, stdin, cwd, env, timeout):
            seen.update(cmd=cmd, stdin=stdin, env=env, system=Path(cmd[cmd.index("--system-prompt-file") + 1]).read_text(encoding="utf-8"))
            out = {"type": "result", "subtype": "success", "is_error": False, "result": "Dạ được ạ",
                   "usage": {"input_tokens": 12, "output_tokens": 5, "cache_read_input_tokens": 900, "cache_creation_input_tokens": 40},
                   "modelUsage": {"claude-sonnet-5-5": {}}}
            return subprocess.CompletedProcess(cmd, 0, json.dumps(out).encode(), b"")
        with mock.patch.object(cli_llm, "_run", fake_run):
            r = cli_llm.run("claude-cli", "Hệ thống: trả lời ngắn", "Xin chào")
        self.assertEqual((r["text"], r["model"], r["input"], r["output"], r["cache_read"]), ("Dạ được ạ", "claude-sonnet-5-5", 12, 5, 900))
        cmd = seen["cmd"]
        self.assertEqual(cmd[:2], ["/bin/claude", "-p"])
        self.assertEqual(cmd[cmd.index("--tools") + 1], "")
        self.assertEqual(cmd[cmd.index("--setting-sources") + 1], "")
        self.assertIn("--strict-mcp-config", cmd)
        self.assertEqual(seen["stdin"], "Xin chào")
        self.assertEqual(seen["system"], "Hệ thống: trả lời ngắn")
        self.assertNotIn("ANTHROPIC_API_KEY", seen["env"])                   # không lẫn sang API key

    def test_codex_reads_last_message_and_usage(self):
        def fake_run(cmd, stdin, cwd, env, timeout):
            Path(cmd[cmd.index("-o") + 1]).write_text("Kết quả từ ChatGPT", encoding="utf-8")
            self.assertNotIn("OPENAI_API_KEY", env)
            self.assertEqual(cmd[-1], "-")
            self.assertIn("# Hướng dẫn\nsys", stdin)
            lines = [json.dumps({"type": "thread.started"}),
                     json.dumps({"type": "turn.completed", "usage": {"input_tokens": 1000, "cached_input_tokens": 600, "output_tokens": 50}})]
            return subprocess.CompletedProcess(cmd, 0, "\n".join(lines).encode(), b"")
        with mock.patch.object(cli_llm, "_run", fake_run):
            r = cli_llm.run("codex-cli", "sys", "câu hỏi")
        self.assertEqual((r["text"], r["input"], r["cache_read"], r["output"]), ("Kết quả từ ChatGPT", 400, 600, 50))

    def test_gemini_json_after_log_lines(self):
        def fake_run(cmd, stdin, cwd, env, timeout):
            self.assertEqual(Path(env["GEMINI_SYSTEM_MD"]).read_text(encoding="utf-8"), "sys")
            body = json.dumps({"response": "Trả lời từ Gemini", "stats": {"models": {"gemini-3-pro": {"tokens": {"prompt": 300, "candidates": 20, "cached": 100, "thoughts": 5}}}}}, indent=2)
            return subprocess.CompletedProcess(cmd, 0, ("Loaded cached credentials.\n" + body).encode(), b"")
        with mock.patch.object(cli_llm, "_run", fake_run):
            r = cli_llm.run("gemini-cli", "sys", "câu hỏi")
        self.assertEqual((r["text"], r["model"], r["input"], r["output"], r["cache_read"]), ("Trả lời từ Gemini", "gemini-3-pro", 200, 25, 100))

    def test_not_logged_in_gives_login_hint(self):
        def fake_run(cmd, stdin, cwd, env, timeout):
            return subprocess.CompletedProcess(cmd, 1, b"", b"Error: Not logged in. Please run codex login")
        with mock.patch.object(cli_llm, "_run", fake_run), self.assertRaises(RuntimeError) as cm:
            cli_llm.run("codex-cli", "sys", "x")
        self.assertIn("codex login", str(cm.exception))

    def test_disabled_in_tests_by_default(self):
        with mock.patch.dict(os.environ, {"CLI_LLM_DISABLED": "1"}), self.assertRaises(RuntimeError):
            cli_llm.run("claude-cli", "s", "p")

    @unittest.skipUnless(os.name == "nt", "shim .cmd của npm chỉ có trên Windows")
    def test_npm_cmd_shim_runs_node_directly(self):
        self.which.stop()
        try:
            with tempfile.TemporaryDirectory() as d:
                js = Path(d) / "node_modules" / "@x" / "cli" / "bin" / "x.js"
                js.parent.mkdir(parents=True)
                js.write_text("//", encoding="utf-8")
                (Path(d) / "node.exe").write_bytes(b"")
                shim = Path(d) / "xcli.cmd"
                shim.write_text('@ECHO off\r\n"%_prog%"  "%dp0%\\node_modules\\@x\\cli\\bin\\x.js" %*\r\n', encoding="utf-8")
                cli_llm._which_cache.clear()
                with mock.patch.object(cli_llm.shutil, "which", lambda name: str(shim) if name == "xcli" else None):
                    self.assertEqual(cli_llm.resolve_cmd("xcli"), [str(Path(d) / "node.exe"), str(js)])
                cli_llm._which_cache.clear()
        finally:
            self.which.start()


class GatewaySubscriptionTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        reset_db()
        self.mid = db.create_meeting("Họp gói đăng ký")

    async def test_cli_provider_is_used_and_logged_free(self):
        artifacts.set_provider("claude-cli", fallback=False, model="opus")
        seen = {}

        def fake(provider, system, prompt, timeout=None, model=None):
            seen.update(provider=provider, model=model, system=system)
            return {"text": "ok qua gói", "model": "claude-opus-5-5", "input": 10, "output": 3, "cache_read": 0,
                    "cache_write": 0, "estimated": False}
        with mock.patch.object(cli_llm, "run", fake), mock.patch.object(cli_llm, "installed", lambda p: True):
            self.assertTrue(artifacts.llm_available())
            artifacts.set_meeting(self.mid, "slide")
            sysb = artifacts.Blocks([artifacts.text_block("phần đầu", cache=True), artifacts.text_block("phần sau")])
            self.assertEqual(await artifacts._call_llm(sysb, "hỏi"), "ok qua gói")
        self.assertEqual((seen["provider"], seen["model"], seen["system"]), ("claude-cli", "opus", "phần đầu\n\nphần sau"))
        row = db._get_db()["llm_usage"].find_one({"meeting_id": self.mid})
        self.assertEqual((row["provider"], row["model"], row["cost_usd"]), ("claude-cli", "claude-opus-5-5", 0.0))

    async def test_fallback_to_api_only_when_enabled(self):
        def boom(*a, **k):
            raise RuntimeError("hết hạn mức")
        msgs = SimpleNamespace(calls=[])

        async def create(**kw):
            msgs.calls.append(kw)
            return SimpleNamespace(content=[SimpleNamespace(type="text", text="qua API")],
                                   usage=SimpleNamespace(input_tokens=5, output_tokens=2, cache_read_input_tokens=0))
        client = SimpleNamespace(messages=SimpleNamespace(create=create))
        with mock.patch.object(cli_llm, "run", boom), mock.patch.object(artifacts, "_anthropic", lambda: client), \
                mock.patch.dict(os.environ, {"ANTHROPIC_API_KEY": "x"}):
            artifacts.set_provider("codex-cli", fallback=False)
            with self.assertRaises(RuntimeError):
                await artifacts._call_llm("s", "p")
            artifacts.set_provider("codex-cli", fallback=True)
            self.assertEqual(await artifacts._call_llm("s", "p"), "qua API")
        self.assertEqual(len(msgs.calls), 1)

    async def test_invalid_provider_and_model_refused(self):
        with self.assertRaises(ValueError):
            artifacts.set_provider("chatgpt-web")
        with self.assertRaises(ValueError):
            artifacts.set_provider("claude-cli", model="sonnet; rm -rf /")


class DocumentApiTests(TempLibrary, unittest.TestCase):
    def setUp(self):
        TempLibrary.setUp(self)
        reset_db()
        self.client = TestClient(appmod.app)
        self.client.__enter__()
        self.mid = self.client.post("/api/meetings", json={"title": "Họp tài liệu"}).json()["meeting_id"]
        make_pdf(self.base / "bao-cao.pdf", [["Bao cao Q3"], ["Ket luan"]])

    def tearDown(self):
        self.client.__exit__(None, None, None)
        live.SESSIONS.clear()
        TempLibrary.tearDown(self)

    def test_library_open_file_assets_and_present_mode(self):
        lib = self.client.get("/api/library").json()
        self.assertEqual([it["name"] for it in lib["items"]], ["bao-cao"])
        path = lib["items"][0]["path"]
        bad = self.client.post(f"/api/meetings/{self.mid}/open-file", json={"path": str(Path(self.tmp.name) / "x.pdf")})
        self.assertEqual(bad.status_code, 400)
        r = self.client.post(f"/api/meetings/{self.mid}/open-file", json={"path": path}).json()
        ask = [e for e in r["events"] if e["type"] == "ai_ask"]
        self.assertEqual(ask[0]["kind"], "present_mode")
        art = self.client.get(f"/api/artifacts/{r['artifact_id']}").json()
        img = json.loads(art["content"])["slides"][0]["image"]
        got = self.client.get(img)
        self.assertEqual((got.status_code, got.headers["content-type"]), (200, "image/jpeg"))
        self.assertEqual(self.client.get("/api/deck-assets/zzzz/1.jpg").status_code, 404)
        res = self.client.post(f"/api/meetings/{self.mid}/present-mode", json={"mode": "human", "artifact_id": r["artifact_id"]}).json()
        self.assertEqual(res["mode"], "human")
        self.assertEqual(self.client.post(f"/api/meetings/{self.mid}/present-mode", json={"mode": "script", "source": "x"}).status_code, 400)

    def test_library_dirs_setting(self):
        extra = Path(self.tmp.name) / "Tài liệu"
        extra.mkdir()
        (extra / "kich-ban.docx").write_bytes(b"")
        r = self.client.put("/api/settings/library", json={"dirs": [str(extra)]}).json()
        self.assertEqual(r["saved"], [str(extra)])
        self.assertIn(str(extra), [x["path"] for x in r["roots"]])
        self.assertEqual(self.client.put("/api/settings/library", json={"dirs": [str(extra / "khong-co")]}).status_code, 400)
        names = [it["name"] for it in self.client.get("/api/library?q=kich ban").json()["items"]]
        self.assertEqual(names, ["kich-ban"])

    def test_provider_endpoints(self):
        info = self.client.get("/api/llm/providers").json()
        self.assertEqual([p["id"] for p in info["providers"]], ["claude", "gemini", "claude-cli", "codex-cli", "gemini-cli"])
        self.assertEqual(info["current"], "claude")
        self.assertEqual(self.client.put("/api/llm/provider", json={"provider": "abc"}).status_code, 400)
        r = self.client.put("/api/llm/provider", json={"provider": "gemini-cli", "api_fallback": True, "model": "gemini-3-pro"}).json()
        self.assertEqual((r["current"], r["api_fallback"]), ("gemini-cli", True))
        self.assertEqual(next(p for p in r["providers"] if p["id"] == "gemini-cli")["model_setting"], "gemini-3-pro")

        def fake(provider, system, prompt, timeout=None, model=None):
            return {"text": "Xin chào!", "model": model, "input": 5, "output": 3, "cache_read": 0, "cache_write": 0, "estimated": False}
        with mock.patch.object(cli_llm, "run", fake):
            t = self.client.post("/api/llm/test", json={"provider": "gemini-cli"}).json()
        self.assertEqual((t["ok"], t["text"]), (True, "Xin chào!"))
        self.assertEqual(self.client.get("/api/health").json()["llm"]["provider"], "gemini-cli")


if __name__ == "__main__":
    unittest.main()
